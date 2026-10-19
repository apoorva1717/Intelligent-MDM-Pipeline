"""tools/vat_validate_xlsx.py on the VAT test export (tests/fixtures/vat/).

``VAT_TEST_EXPORT_30.xlsx`` is the hand-built fixture: sheet ``Data`` in the
SAP export layout (28 columns + VAT Registration No.), sheet ``Test_Notes``
with the expected outcome per row. VIES is mocked — every FORMAT_OK number
answers VALID-undisclosed — so these tests pin the offline half (format
status, flags, fan-out, conflict) against Test_Notes, not VIES's verdicts.
"""

from __future__ import annotations

import hashlib
import logging
from dataclasses import replace
from pathlib import Path

import pytest
import pytest_asyncio
from openpyxl import Workbook, load_workbook

from config import get_settings
from tests.mocks.vies_mock import ViesStub, echo_valid_undisclosed, error
from tools import vat_validate_xlsx
from utils.cache import EvidenceCache
from vat.validate import CUSTOMER_VAT_CONFLICT, VAT_COLUMNS
from vat.vies_client import VIESClient

FIXTURE = Path(__file__).parent / "fixtures" / "vat" / "VAT_TEST_EXPORT_30.xlsx"

#: Test_Notes "Offline format check" → VAT_Format_Status.
_NOTE_TO_STATUS = {"pass": {"FORMAT_OK"}, "FAIL": {"FORMAT_ERROR"},
                   "(blank)": {"MISSING", "NOT_EU"}}


def _client(tmp_path, stub):
    settings = replace(get_settings(), vies_backoff_seconds=0.0, vies_min_interval_ms=0,
                       vies_requester_vat="", cache_frozen=False)
    return VIESClient(settings, cache=EvidenceCache(tmp_path / "cache"),
                      transport=stub.transport)


def _sheet_rows(path, sheet):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    return list(rows[0]), [dict(zip(rows[0], r)) for r in rows[1:]]


@pytest_asyncio.fixture
async def ran(tmp_path):
    stub = ViesStub(echo_valid_undisclosed)
    out = tmp_path / "out.xlsx"
    summary = await vat_validate_xlsx.run(FIXTURE, out, _client(tmp_path, stub))
    return out, summary, stub


class TestAgainstTestNotes:
    @pytest.mark.asyncio
    async def test_every_row_matches_its_offline_expectation(self, ran):
        out, _, _ = ran
        _, data = _sheet_rows(out, "Data")
        _, notes = _sheet_rows(FIXTURE, "Test_Notes")
        assert len(data) == len(notes) == 30
        for row, note in zip(data, notes):
            assert row["Customer"] == note["Customer"]
            assert row["VAT_Format_Status"] in _NOTE_TO_STATUS[note["Offline format check"]], note

    @pytest.mark.asyncio
    async def test_named_cases(self, ran):
        out, _, _ = ran
        _, data = _sheet_rows(out, "Data")
        by_row = {i: r for i, r in enumerate(data, start=2)}  # sheet row numbers

        bosch = [by_row[4], by_row[5]]  # customer split → validated once, fanned out
        assert {r["VAT_VIES_Status"] for r in bosch} == {"VALID"}
        assert all(r["VAT_Flags"] == "NAME_MATCH:NOT_DISCLOSED" for r in bosch)

        bayer = [by_row[8], by_row[9]]  # conflicting VAT (typo) → flagged, none picked
        assert all(CUSTOMER_VAT_CONFLICT in r["VAT_Flags"] for r in bayer)
        assert (by_row[8]["VAT_VIES_Status"], by_row[9]["VAT_VIES_Status"]) == ("VALID", "NOT_CHECKED")

        assert by_row[16]["VAT_Normalized"] == "DE147799625"       # prefix added
        assert by_row[16]["VAT_Flags"] == "PREFIX_FROM_COUNTRY_KEY; NAME_MATCH:NOT_DISCLOSED"
        assert by_row[19]["VAT_VIES_Status"] == "NOT_CHECKED"      # bad checksum, no call
        assert by_row[19]["VAT_Flags"] == "FORMAT_ERROR:InvalidChecksum; NAME_MATCH:VIES_NOT_VALID"
        assert by_row[20]["VAT_Format_Status"] == "MISSING"
        assert by_row[21]["VAT_Flags"].startswith("PREFIX_COUNTRY_MISMATCH")  # IE prefix, DE key
        assert by_row[31]["VAT_Format_Status"] == "NOT_EU"         # US customer
        assert by_row[12]["VAT_Normalized"] == "DE119429301"       # "DE 11 94 29 301"
        assert by_row[11]["VAT_Normalized"] == "DE811850788"       # lowercase prefix

    @pytest.mark.asyncio
    async def test_vies_is_asked_once_per_unique_number(self, ran):
        out, summary, stub = ran
        _, data = _sheet_rows(out, "Data")
        eligible = {r["VAT_Normalized"] for r in data if r["VAT_Format_Status"] == "FORMAT_OK"}
        assert len(stub.requests) == len(eligible) == summary["unique_vies_lookups"] == 24  # 26 FORMAT_OK rows: Bosch x2, IE6388047V x2


class TestWorkbook:
    @pytest.mark.asyncio
    async def test_input_columns_and_other_sheets_unchanged(self, ran):
        out, _, _ = ran
        in_headers, in_rows = _sheet_rows(FIXTURE, "Data")
        out_headers, out_rows = _sheet_rows(out, "Data")
        assert out_headers == in_headers + VAT_COLUMNS
        assert [{h: r[h] for h in in_headers} for r in out_rows] == in_rows
        assert _sheet_rows(out, "Test_Notes") == _sheet_rows(FIXTURE, "Test_Notes")
        assert load_workbook(out).sheetnames == ["Data", "Test_Notes", "Summary"]

    @pytest.mark.asyncio
    async def test_input_file_is_not_modified(self, tmp_path):
        before = hashlib.sha1(FIXTURE.read_bytes()).hexdigest()
        await vat_validate_xlsx.run(FIXTURE, tmp_path / "o.xlsx",
                                    _client(tmp_path, ViesStub(echo_valid_undisclosed)))
        assert hashlib.sha1(FIXTURE.read_bytes()).hexdigest() == before

    @pytest.mark.asyncio
    async def test_summary_sheet(self, ran):
        out, summary, _ = ran
        ws = load_workbook(out, read_only=True)["Summary"]
        cells = {r[0]: r[1:] for r in ws.iter_rows(values_only=True) if r and r[0]}
        assert cells["Rows"][0] == 30
        assert cells["Unique customers"][0] == 28
        assert cells["Customers with conflicting VAT values"][0] == 1
        assert cells["FORMAT_OK"][0] == 26
        assert cells["NOT_AVAILABLE"][0] == 30  # mock VIES discloses no names
        assert cells["DE"][0] == 20          # Country/Region Key DE: 20 rows
        assert cells["US"][0] == 1
        assert summary["rows"] == 30

    @pytest.mark.asyncio
    async def test_rerun_on_output_overwrites_vat_columns(self, ran, tmp_path):
        out, _, _ = ran
        again = tmp_path / "again.xlsx"
        await vat_validate_xlsx.run(out, again, _client(tmp_path, ViesStub(echo_valid_undisclosed)))
        headers, _ = _sheet_rows(again, "Data")
        assert headers == _sheet_rows(out, "Data")[0]
        assert load_workbook(again).sheetnames.count("Summary") == 1

    @pytest.mark.asyncio
    async def test_falls_back_to_active_sheet_with_a_warning(self, tmp_path, caplog):
        src = tmp_path / "no_data_sheet.xlsx"
        wb = Workbook()
        wb.active.title = "Export"
        wb.active.append(["Customer", "Country/Region Key", "VAT Registration No."])
        wb.active.append(["1", "DE", "DE136695976"])
        wb.save(src)
        with caplog.at_level(logging.WARNING, logger="vat.xlsx"):
            await vat_validate_xlsx.run(src, tmp_path / "o.xlsx",
                                        _client(tmp_path, ViesStub(echo_valid_undisclosed)))
        assert "No sheet named 'Data'" in caplog.text
        headers, rows = _sheet_rows(tmp_path / "o.xlsx", "Export")
        assert rows[0]["VAT_VIES_Status"] == "VALID"


class TestRerunAfterUnavailable:
    @pytest.mark.asyncio
    async def test_summary_lists_unanswered_ids_and_a_rerun_asks_only_for_them(self, tmp_path):
        down = {"IE6388047V", "SE556703748501"}

        def reply(body):
            vat = body["countryCode"] + body["vatNumber"]
            return error("MS_UNAVAILABLE") if vat in down else echo_valid_undisclosed(body)

        first = tmp_path / "first.xlsx"
        summary = await vat_validate_xlsx.run(FIXTURE, first, _client(tmp_path, ViesStub(reply)))
        assert summary["vies_service_unavailable_vat_ids"] == 2
        assert summary["vies_service_unavailable_rows"] == 3  # IE6388047V is on two rows
        assert [(u["vat"], u["error_code"], u["rows"]) for u in summary["vies_unanswered"]] == [
            ("IE6388047V", "MS_UNAVAILABLE", 2), ("SE556703748501", "MS_UNAVAILABLE", 1),
        ]

        ws = load_workbook(first, read_only=True)["Summary"]
        cells = {r[0]: r[1:] for r in ws.iter_rows(values_only=True) if r and r[0]}
        assert cells["Unique VAT IDs SERVICE_UNAVAILABLE"][0] == 2
        assert cells["Rows SERVICE_UNAVAILABLE"][0] == 3
        assert cells["IE6388047V"][:3] == ("SERVICE_UNAVAILABLE", "MS_UNAVAILABLE", 2)

        # Same input, same cache: only the two unanswered numbers go to VIES.
        stub = ViesStub(echo_valid_undisclosed)
        again = await vat_validate_xlsx.run(FIXTURE, tmp_path / "second.xlsx", _client(tmp_path, stub))
        assert sorted(r["countryCode"] + r["vatNumber"] for r in stub.requests) == sorted(down)
        assert again["vies_from_cache"] == 22
        assert again["vies_service_unavailable_vat_ids"] == 0
