"""POST /api/validate/vat (api/vat_routes.py + vat/validate.py), VIES mocked."""

from __future__ import annotations

import io
from dataclasses import replace
from pathlib import Path

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook

from api import vat_routes
from api.app import app
from config import get_settings
from tests.mocks import vies_mock
from tests.mocks.vies_mock import ViesStub
from utils.cache import EvidenceCache
from vat.validate import CUSTOMER_VAT_CONFLICT, VAT_COLUMNS
from vat.vies_client import VIESClient

URL = "/api/validate/vat"

#: The 28 SAP export columns + VAT Registration No., in export order.
EXPORT_HEADERS = [
    "Customer", "Central Deletion Flag", "Central delivery block", "Company Code",
    "Sales Organization", "Distribution Channel", "Division", "Name", "Name 2",
    "Name 3", "Name 4", "PO Box", "House Number", "Street", "Street 2", "Street 3",
    "Street 4", "City", "Postal Code", "Region", "Country/Region Key",
    "Account group", "Reconciliation acct", "Comments", "Terms of Payment",
    "Sales District", "Created On", "Created By", "VAT Registration No.",
]


def _row(customer, country, vat, **extra):
    row = dict.fromkeys(EXPORT_HEADERS)
    row.update({
        "Customer": customer, "Country/Region Key": country,
        "VAT Registration No.": vat, "Name": f"Example {customer}",
        "Created By": "VATTEST",
    })
    row.update(extra)
    return row


@pytest.fixture
def stub():
    return ViesStub(vies_mock.echo_valid_undisclosed)


@pytest_asyncio.fixture
async def client(stub, tmp_path, monkeypatch):
    settings = replace(
        get_settings(), vies_backoff_seconds=0.0, vies_min_interval_ms=0,
        vies_requester_vat="", cache_frozen=False,
    )

    def factory(_settings=None):
        return VIESClient(settings, cache=EvidenceCache(tmp_path), transport=stub.transport)

    monkeypatch.setattr(vat_routes, "_get_vies_client", factory)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


async def _post(client, rows):
    resp = await client.post(URL, json={"rows": rows})
    assert resp.status_code == 200, resp.text
    return resp.json()


class TestOutputShape:
    @pytest.mark.asyncio
    async def test_input_columns_unchanged_and_vat_columns_appended(self, client):
        row = _row("99000001", "DE", "DE 136.695-976", **{"Company Code": 1000, "Street": " x "})
        body = await _post(client, [row])
        out = body["rows"][0]
        assert list(out) == EXPORT_HEADERS + VAT_COLUMNS
        assert {k: out[k] for k in EXPORT_HEADERS} == row
        assert body["columns"] == VAT_COLUMNS

    @pytest.mark.asyncio
    async def test_de_valid_with_undisclosed_name(self, client):
        out = (await _post(client, [_row("1", "DE", "DE136695976")]))["rows"][0]
        assert out["VAT_Normalized"] == "DE136695976"
        assert out["VAT_Format_Status"] == "FORMAT_OK"
        assert out["VAT_VIES_Status"] == "VALID"
        assert (out["VAT_VIES_Name"], out["VAT_VIES_Address"]) == (None, None)
        assert out["VAT_Name_Disclosed"] is False
        assert out["VAT_Request_Date"] == vies_mock.REQUEST_DATE
        assert out["VAT_Consultation_No"] is None
        assert out["VAT_Name_Match"] == "NOT_AVAILABLE"
        assert out["VAT_Flags"] == "NAME_MATCH:NOT_DISCLOSED"

    @pytest.mark.asyncio
    async def test_headers_are_matched_tolerantly(self, client):
        rows = [{"customer": "1", "country_region_key": "DE", "VAT REGISTRATION NO": "136695976"}]
        out = (await _post(client, rows))["rows"][0]
        assert (out["VAT_Normalized"], out["VAT_Flags"]) == (
            "DE136695976", "PREFIX_FROM_COUNTRY_KEY; NAME_MATCH:NOT_DISCLOSED",
        )


class TestFanOut:
    @pytest.mark.asyncio
    async def test_customer_split_is_checked_once(self, client, stub):
        rows = [
            _row("99000003", "DE", "DE136695976", **{"Company Code": "1000"}),
            _row("99000003", "DE", "DE136695976", **{"Company Code": "1100"}),
        ]
        body = await _post(client, rows)
        assert [r["VAT_VIES_Status"] for r in body["rows"]] == ["VALID", "VALID"]
        assert len(stub.requests) == 1
        assert body["summary"]["unique_vies_lookups"] == 1

    @pytest.mark.asyncio
    async def test_same_number_spelled_differently_is_not_a_conflict(self, client, stub):
        rows = [_row("7", "DE", "DE136695976"), _row("7", "DE", "de 136 695 976")]
        body = await _post(client, rows)
        assert all(CUSTOMER_VAT_CONFLICT not in r["VAT_Flags"] for r in body["rows"])
        assert len(stub.requests) == 1

    @pytest.mark.asyncio
    async def test_conflicting_values_are_flagged_and_neither_is_picked(self, client):
        rows = [_row("99000006", "DE", "DE136695976"), _row("99000006", "DE", "DE136695977")]
        body = await _post(client, rows)
        a, b = body["rows"]
        assert (a["VAT_Normalized"], a["VAT_VIES_Status"]) == ("DE136695976", "VALID")
        assert (b["VAT_Normalized"], b["VAT_Format_Status"], b["VAT_VIES_Status"]) == (
            "DE136695977", "FORMAT_ERROR", "NOT_CHECKED",
        )
        assert CUSTOMER_VAT_CONFLICT in a["VAT_Flags"]
        assert CUSTOMER_VAT_CONFLICT in b["VAT_Flags"]
        assert body["summary"]["customers_with_vat_conflict"] == 1

    @pytest.mark.asyncio
    async def test_value_on_one_row_and_blank_on_another_is_a_conflict(self, client):
        rows = [_row("8", "DE", "DE136695976"), _row("8", "DE", None)]
        body = await _post(client, rows)
        assert all(CUSTOMER_VAT_CONFLICT in r["VAT_Flags"] for r in body["rows"])

    @pytest.mark.asyncio
    async def test_customer_is_grouped_as_read_no_zero_stripping(self, client):
        rows = [_row("0099000006", "DE", "DE136695976"), _row("99000006", "DE", None)]
        body = await _post(client, rows)
        assert all(r["VAT_Flags"] is None or CUSTOMER_VAT_CONFLICT not in r["VAT_Flags"]
                   for r in body["rows"])
        assert body["summary"]["unique_customers"] == 2


class TestStates:
    @pytest.mark.asyncio
    async def test_format_error_is_never_sent(self, client, stub):
        out = (await _post(client, [_row("1", "DE", "DE136695977")]))["rows"][0]
        assert (out["VAT_Format_Status"], out["VAT_VIES_Status"]) == ("FORMAT_ERROR", "NOT_CHECKED")
        assert out["VAT_Flags"] == "FORMAT_ERROR:InvalidChecksum; NAME_MATCH:VIES_NOT_VALID"
        assert stub.requests == []

    @pytest.mark.asyncio
    async def test_us_row_and_missing_row(self, client, stub):
        body = await _post(client, [_row("1", "US", None), _row("2", "DE", None)])
        us, de = body["rows"]
        assert (us["VAT_Format_Status"], us["VAT_VIES_Status"]) == ("NOT_EU", "NOT_CHECKED")
        assert (de["VAT_Format_Status"], de["VAT_VIES_Status"]) == ("MISSING", "NOT_CHECKED")
        assert stub.requests == []

    @pytest.mark.asyncio
    async def test_prefix_country_mismatch_is_flagged(self, client):
        out = (await _post(client, [_row("1", "DE", "FR40303265045")]))["rows"][0]
        assert (out["VAT_Normalized"], out["VAT_VIES_Status"]) == ("FR40303265045", "VALID")
        assert out["VAT_Flags"] == "PREFIX_COUNTRY_MISMATCH; NAME_MATCH:NOT_DISCLOSED"

    @pytest.mark.asyncio
    @pytest.mark.parametrize("code", ["MS_UNAVAILABLE", "GLOBAL_MAX_CONCURRENT_REQ"])
    async def test_service_unavailable_is_not_invalid(self, client, stub, code):
        stub.replies = [vies_mock.error(code)]
        out = (await _post(client, [_row("1", "DE", "DE136695976")]))["rows"][0]
        assert out["VAT_VIES_Status"] == "SERVICE_UNAVAILABLE"
        assert out["VAT_Flags"] == f"VIES:{code}; NAME_MATCH:VIES_NOT_VALID"
        assert out["VAT_Name_Match"] == "NOT_AVAILABLE"

    @pytest.mark.asyncio
    async def test_disclosed_name_and_address(self, client, stub):
        stub.replies = [vies_mock.valid_disclosed("FR", "40303265045")]
        out = (await _post(client, [_row("1", "FR", "FR40303265045")]))["rows"][0]
        assert (out["VAT_VIES_Name"], out["VAT_Name_Disclosed"]) == ("EXAMPLE TRADING SA", True)
        assert (out["VAT_Name_Match"], out["VAT_Flags"]) == ("MISMATCH", "NAME_MATCH:NO_MATCH")

    @pytest.mark.asyncio
    async def test_name_match_reads_name_and_name_2_to_4(self, client, stub):
        stub.replies = [vies_mock.valid_disclosed("FR", "40303265045", name="SA EXAMPLE TRADING")]
        row = _row("1", "FR", "FR40303265045", Name="Example", **{"Name 2": "Trading S.A."})
        out = (await _post(client, [row]))["rows"][0]
        assert (out["VAT_Name_Match"], out["VAT_Flags"]) == ("MATCH", "NAME_MATCH:EXACT_WITH_NAME_2_4")


class TestSummaryAndErrors:
    @pytest.mark.asyncio
    async def test_summary_counts(self, client):
        rows = [
            _row("1", "DE", "DE136695976"), _row("1", "DE", "DE136695976"),
            _row("2", "DE", "DE136695977"), _row("3", "US", None),
        ]
        s = (await _post(client, rows))["summary"]
        assert (s["rows"], s["unique_customers"], s["unique_vies_lookups"]) == (4, 3, 1)
        assert s["format_status"] == {"FORMAT_OK": 2, "FORMAT_ERROR": 1, "NOT_EU": 1}
        assert s["vies_status"] == {"VALID": 2, "NOT_CHECKED": 2}
        assert s["by_country"]["US"] == {
            "format_status": {"NOT_EU": 1}, "vies_status": {"NOT_CHECKED": 1},
        }

    @pytest.mark.asyncio
    async def test_no_vat_column_is_422(self, client):
        resp = await client.post(URL, json={"rows": [{"Customer": "1", "Country/Region Key": "DE"}]})
        assert resp.status_code == 422

    @pytest.mark.asyncio
    async def test_empty_batch_is_422(self, client):
        resp = await client.post(URL, json={"rows": []})
        assert resp.status_code == 422



# ---------------------------------------------------------------------------
# POST /api/validate/vat/file
# ---------------------------------------------------------------------------

FILE_URL = "/api/validate/vat/file"
FIXTURE = Path(__file__).parent / "fixtures" / "vat" / "VAT_TEST_EXPORT_30.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _upload(client, name, content):
    return await client.post(FILE_URL, files={"file": (name, content, XLSX)})


def _book(rows):
    wb = Workbook()
    wb.active.title = "Data"
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(wb, title):
    rows = list(wb[title].iter_rows(values_only=True))
    return list(rows[0]), [dict(zip(rows[0], r)) for r in rows[1:]]


class TestFileEndpoint:
    @pytest.mark.asyncio
    async def test_fixture_round_trip(self, client, stub):
        resp = await _upload(client, "VAT_TEST_EXPORT_30.xlsx", FIXTURE.read_bytes())
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == XLSX
        assert resp.headers["content-disposition"] == 'attachment; filename="VAT_TEST_EXPORT_30_vat.xlsx"'
        assert (resp.headers["x-vat-rows"], resp.headers["x-vat-service-unavailable"]) == ("30", "0")

        out = load_workbook(io.BytesIO(resp.content), read_only=True)
        src = load_workbook(FIXTURE, read_only=True)
        assert out.sheetnames == ["Data", "Test_Notes", "Summary"]
        in_headers, in_rows = _sheet(src, "Data")
        out_headers, out_rows = _sheet(out, "Data")
        assert out_headers == in_headers + VAT_COLUMNS
        assert [{h: r[h] for h in in_headers} for r in out_rows] == in_rows
        assert _sheet(out, "Test_Notes") == _sheet(src, "Test_Notes")
        assert [r["VAT_Format_Status"] for r in out_rows].count("FORMAT_OK") == 26
        assert len(stub.requests) == 24  # once per unique VAT, not per row

    @pytest.mark.asyncio
    async def test_service_unavailable_count_header(self, client, stub):
        stub.replies = [vies_mock.error("MS_UNAVAILABLE")]
        content = _book([["Customer", "Country/Region Key", "VAT Registration No."],
                         ["1", "DE", "DE136695976"], ["2", "FR", "FR40303265045"]])
        resp = await _upload(client, "x.xlsx", content)
        assert resp.headers["x-vat-service-unavailable"] == "2"

    @pytest.mark.asyncio
    async def test_active_sheet_fallback(self, client):
        wb = Workbook()
        wb.active.title = "Export"
        wb.active.append(["Customer", "Country/Region Key", "VAT Registration No."])
        wb.active.append(["1", "DE", "DE136695976"])
        buf = io.BytesIO()
        wb.save(buf)
        resp = await _upload(client, "export.xlsx", buf.getvalue())
        assert resp.status_code == 200
        _, rows = _sheet(load_workbook(io.BytesIO(resp.content)), "Export")
        assert rows[0]["VAT_VIES_Status"] == "VALID"

    @pytest.mark.asyncio
    @pytest.mark.parametrize("name, content, status", [
        ("export.csv", b"Customer,VAT", 400),   # wrong extension
        ("export.xlsx", b"", 400),                       # empty
        ("export.xlsx", b"not a workbook", 400),         # unreadable
    ])
    async def test_bad_uploads(self, client, name, content, status):
        resp = await _upload(client, name, content)
        assert resp.status_code == status

    @pytest.mark.asyncio
    async def test_no_vat_column_is_422(self, client):
        resp = await _upload(client, "x.xlsx", _book([["Customer", "Name"], ["1", "Acme"]]))
        assert resp.status_code == 422
        assert "VAT Registration No." in resp.json()["detail"]

    @pytest.mark.asyncio
    async def test_no_data_rows_is_422(self, client):
        resp = await _upload(client, "x.xlsx", _book([["Customer", "VAT Registration No."]]))
        assert resp.status_code == 422



# ---------------------------------------------------------------------------
# lookup_missing (none | grounded | llm_web | llm_only) and the experiment alias.
# The endpoints never answer from memory: llm_only runs llm_web.
# ---------------------------------------------------------------------------

from vat.lookup import COMPARE_COLUMNS, LOOKUP_COLUMNS, FoundResult  # noqa: E402


class _FakeLookup:
    """Proposes one fixed value for every in-scope customer."""

    def __init__(self, mode):
        self.mode = mode
        self.calls = 0
        self.closed = False

    async def lookup(self, inp):
        self.calls += 1
        status = "FOUND_VERIFIED" if self.mode == "grounded" else "LLM_GUESS:FORMAT_OK:VALID"
        return FoundResult(self.mode, status, value="DE136695976")

    def stats(self):
        return {"calls": self.calls}

    async def aclose(self):
        self.closed = True


@pytest.fixture
def lookups(monkeypatch):
    made: list = []

    def factory(mode, client):
        if mode == "none":
            return None
        engine = _FakeLookup(mode)
        made.append(engine)
        return engine

    monkeypatch.setattr(vat_routes, "_get_lookup", factory)
    return made


_MISSING_DE = {"Customer": "1", "Name": "Acme GmbH", "Country/Region Key": "DE",
               "VAT Registration No.": None}


class TestLookupMissing:
    @pytest.mark.asyncio
    async def test_default_is_none_and_output_unchanged(self, client, lookups):
        body = await _post(client, [_MISSING_DE])
        assert body["columns"] == VAT_COLUMNS
        assert list(body["rows"][0]) == list(_MISSING_DE) + VAT_COLUMNS
        assert "lookup" not in body["summary"] and lookups == []

    @pytest.mark.asyncio
    async def test_body_parameter(self, client, lookups):
        resp = await client.post(URL, json={"rows": [_MISSING_DE], "lookup_missing": "grounded"})
        body = resp.json()
        assert body["columns"] == VAT_COLUMNS + LOOKUP_COLUMNS
        assert body["rows"][0]["VAT_Found_Status"] == "FOUND_VERIFIED"
        assert body["rows"][0]["VAT Registration No."] is None  # never written
        assert lookups[0].closed

    @pytest.mark.asyncio
    async def test_query_parameter_overrides_body(self, client, lookups):
        resp = await client.post(URL + "?lookup_missing=llm_web",
                                 json={"rows": [_MISSING_DE], "lookup_missing": "grounded"})
        assert resp.json()["rows"][0]["VAT_Lookup_Mode"] == "LLM_WEB"

    @pytest.mark.asyncio
    @pytest.mark.parametrize("where", ["query", "body"])
    async def test_llm_only_never_answers_from_memory(self, client, lookups, where):
        url = URL + ("?lookup_missing=llm_only" if where == "query" else "")
        resp = await client.post(url, json={"rows": [_MISSING_DE], "lookup_missing": "llm_only"})
        assert resp.json()["rows"][0]["VAT_Lookup_Mode"] == "LLM_WEB"
        assert [e.mode for e in lookups] == ["llm_web"]  # the memory engine is never built

    @pytest.mark.asyncio
    async def test_experiment_alias_is_llm_web(self, client, lookups):
        resp = await client.post(URL + "/llm-only-experiment", json={"rows": [_MISSING_DE]})
        assert resp.status_code == 200
        row = resp.json()["rows"][0]
        assert row["VAT_Lookup_Mode"] == "LLM_WEB"
        assert [e.mode for e in lookups] == ["llm_web", "grounded"]  # side by side by default

    @pytest.mark.asyncio
    async def test_unknown_mode_is_422(self, client, lookups):
        resp = await client.post(URL, json={"rows": [_MISSING_DE], "lookup_missing": "guess"})
        assert resp.status_code == 422

    @pytest.mark.asyncio
    async def test_file_endpoint_with_lookup(self, client, lookups):
        content = _book([["Customer", "Name", "Country/Region Key", "VAT Registration No."],
                         ["1", "Acme GmbH", "DE", None], ["2", "Beta SA", "FR", "FR40303265045"]])
        resp = await client.post(FILE_URL + "?lookup_missing=grounded",
                                 files={"file": ("x.xlsx", content, XLSX)})
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        headers, rows = _sheet(wb, "Data")
        assert headers[-len(LOOKUP_COLUMNS):] == LOOKUP_COLUMNS
        assert [r["VAT_Found_Status"] for r in rows] == ["FOUND_VERIFIED", None]
        summary_text = [c for r in wb["Summary"].iter_rows(values_only=True) for c in r if c]
        assert any("Missing-VAT lookup" in str(c) for c in summary_text)

    @pytest.mark.asyncio
    async def test_file_endpoint_default_has_no_lookup_columns(self, client, lookups):
        content = _book([["Customer", "Country/Region Key", "VAT Registration No."], ["1", "DE", None]])
        resp = await client.post(FILE_URL, files={"file": ("x.xlsx", content, XLSX)})
        headers, _ = _sheet(load_workbook(io.BytesIO(resp.content)), "Data")
        assert headers[-len(VAT_COLUMNS):] == VAT_COLUMNS and "VAT_Lookup_Mode" not in headers


class TestLlmOnlyExperimentFile:
    @pytest.mark.asyncio
    async def test_file_alias_runs_llm_web(self, client, lookups, stub):
        content = _book([["Customer", "Name", "Country/Region Key", "VAT Registration No."],
                         ["1", "Acme GmbH", "DE", None], ["2", "Beta SA", "FR", "FR40303265045"]])
        resp = await client.post(FILE_URL.replace("/file", "/llm-only-experiment/file"),
                                 files={"file": ("export.xlsx", content, XLSX)})
        # VIES is never called on this endpoint, not even /check-status.
        assert stub.requests == [] and stub.status_requests == 0
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-disposition"] == (
            'attachment; filename="export_vat_llm_web.xlsx"')
        assert resp.headers["x-vat-lookup-mode"] == "llm_web"
        assert [e.mode for e in lookups] == ["llm_web", "grounded"]
        assert all(e.closed for e in lookups)
        wb = load_workbook(io.BytesIO(resp.content))
        _, rows = _sheet(wb, "Data")
        assert [(r["VAT_Lookup_Mode"], r["VAT_Found_Status"]) for r in rows] == [
            ("LLM_WEB", "LLM_GUESS:FORMAT_OK:VALID"), (None, None)]
        assert rows[0]["VAT Registration No."] is None  # never written
        assert rows[1]["VAT_Format_Status"] == "FORMAT_OK"
        assert rows[1]["VAT_VIES_Status"] == "NOT_CHECKED"
        summary_text = [str(c) for r in wb["Summary"].iter_rows(values_only=True) for c in r if c]
        assert any("never from memory" in c for c in summary_text)
        assert not any("EXPERIMENT ONLY" in c for c in summary_text)
        assert "none — VIES not called" in summary_text

    @pytest.mark.asyncio
    async def test_file_alias_rejects_bad_uploads_like_the_file_endpoint(self, client, lookups):
        resp = await client.post(FILE_URL.replace("/file", "/llm-only-experiment/file"),
                                 files={"file": ("export.csv", b"x", XLSX)})
        assert resp.status_code == 400 and lookups == []

    @pytest.mark.asyncio
    async def test_plain_file_endpoint_keeps_its_name_and_mode(self, client, lookups):
        content = _book([["Customer", "Country/Region Key", "VAT Registration No."], ["1", "DE", None]])
        resp = await client.post(FILE_URL, files={"file": ("export.xlsx", content, XLSX)})
        assert resp.headers["content-disposition"] == 'attachment; filename="export_vat.xlsx"'
        assert resp.headers["x-vat-lookup-mode"] == "none"


class TestSideBySide:
    @pytest.mark.asyncio
    async def test_json_has_both_results_and_the_agreement(self, client, lookups):
        resp = await client.post(URL + "/llm-only-experiment", json={"rows": [_MISSING_DE]})
        body = resp.json()
        assert body["columns"] == (VAT_COLUMNS + LOOKUP_COLUMNS + ["VAT_Found_vs_SAP"]
                                   + COMPARE_COLUMNS + ["VAT_Grounded_vs_SAP"])
        row = body["rows"][0]
        assert (row["VAT_Found_Status"], row["VAT_Grounded_Status"]) == (
            "LLM_GUESS:FORMAT_OK:VALID", "FOUND_VERIFIED")
        assert row["VAT_LLM_vs_Grounded"] == "SAME"
        assert body["summary"]["compare_lookup"]["agreement"] == {"SAME": 1}

    @pytest.mark.asyncio
    async def test_with_grounded_false_skips_it(self, client, lookups):
        resp = await client.post(URL + "/llm-only-experiment?with_grounded=false",
                                 json={"rows": [_MISSING_DE]})
        body = resp.json()
        assert body["columns"] == VAT_COLUMNS + LOOKUP_COLUMNS + ["VAT_Found_vs_SAP"]
        assert "VAT_Grounded_Status" not in body["rows"][0]
        assert [e.mode for e in lookups] == ["llm_web"]

    @pytest.mark.asyncio
    async def test_out_of_scope_rows_have_blank_comparison(self, client, lookups):
        fr = {"Customer": "2", "Name": "Beta SA", "Country/Region Key": "FR", "VAT Registration No.": None}
        body = (await client.post(URL + "/llm-only-experiment", json={"rows": [_MISSING_DE, fr]})).json()
        assert body["rows"][1]["VAT_Grounded_Status"] is None
        assert body["rows"][1]["VAT_LLM_vs_Grounded"] is None

    @pytest.mark.asyncio
    async def test_file_has_both_column_blocks(self, client, lookups):
        content = _book([["Customer", "Name", "Country/Region Key", "VAT Registration No."],
                         ["1", "Acme GmbH", "DE", None]])
        resp = await client.post(FILE_URL.replace("/file", "/llm-only-experiment/file"),
                                 files={"file": ("export.xlsx", content, XLSX)})
        wb = load_workbook(io.BytesIO(resp.content))
        headers, rows = _sheet(wb, "Data")
        tail = LOOKUP_COLUMNS + ["VAT_Found_vs_SAP"] + COMPARE_COLUMNS + ["VAT_Grounded_vs_SAP"]
        assert headers[-len(tail):] == tail
        assert rows[0]["VAT_LLM_vs_Grounded"] == "SAME"
        text = [str(c) for r in wb["Summary"].iter_rows(values_only=True) for c in r if c]
        assert any(c.startswith("Side by side") for c in text)

    @pytest.mark.asyncio
    async def test_plain_endpoints_never_add_the_comparison(self, client, lookups):
        body = (await client.post(URL + "?lookup_missing=llm_web", json={"rows": [_MISSING_DE]})).json()
        assert "VAT_LLM_vs_Grounded" not in body["rows"][0]
        assert [e.mode for e in lookups] == ["llm_web"]


class TestAgreement:
    @pytest.mark.parametrize("llm, grounded, expected", [
        ("DE 811 128 135", "DE811128135", "SAME"),
        ("811128135", "de811128135", "SAME"),          # prefix added, case ignored
        ("DE129273398", "DE811128135", "DIFFERENT"),
        ("DE811128135", None, "LLM_ONLY"),
        (None, "DE811128135", "GROUNDED_ONLY"),
        (None, None, "NEITHER"),
    ])
    def test_values(self, llm, grounded, expected):
        from vat.lookup import agreement
        assert agreement(llm, grounded, "DE") == expected



class TestScopeAll:
    @pytest.mark.asyncio
    async def test_experiment_looks_up_rows_that_have_a_vat_id_and_compares_with_sap(self, client, lookups):
        rows = [
            {"Customer": "1", "Name": "Acme GmbH", "Country/Region Key": "DE",
             "VAT Registration No.": "DE136695976"},                     # fake engines propose this
            {"Customer": "2", "Name": "Beta GmbH", "Country/Region Key": "DE",
             "VAT Registration No.": "DE811128135"},
            {"Customer": "3", "Name": "Gamma GmbH", "Country/Region Key": "DE",
             "VAT Registration No.": None},
            {"Customer": "4", "Name": "Delta SA", "Country/Region Key": "FR",
             "VAT Registration No.": "FR40303265045"},
        ]
        body = (await client.post(URL + "/llm-only-experiment", json={"rows": rows})).json()
        out = body["rows"]
        assert [r["VAT_Found_vs_SAP"] for r in out] == ["MATCHES_SAP", "DIFFERS_FROM_SAP", "SAP_BLANK", None]
        assert [r["VAT_Grounded_vs_SAP"] for r in out] == ["MATCHES_SAP", "DIFFERS_FROM_SAP", "SAP_BLANK", None]
        assert [r["VAT Registration No."] for r in out] == [r["VAT Registration No."] for r in rows]
        assert body["summary"]["lookup"]["scope"] == "all"
        assert body["summary"]["lookup"]["rows_in_scope"] == 3       # every DE row
        assert sum(e.calls for e in lookups if e.mode == "llm_web") == 3

    @pytest.mark.asyncio
    async def test_experiment_scope_missing_keeps_the_old_behaviour(self, client, lookups):
        rows = [{"Customer": "1", "Name": "Acme GmbH", "Country/Region Key": "DE",
                 "VAT Registration No.": "DE136695976"}]
        body = (await client.post(URL + "/llm-only-experiment?scope=missing", json={"rows": rows})).json()
        assert body["rows"][0]["VAT_Found_Status"] is None
        assert "VAT_Found_vs_SAP" not in body["rows"][0]

    @pytest.mark.asyncio
    async def test_plain_endpoint_defaults_to_missing(self, client, lookups):
        rows = [{"Customer": "1", "Name": "Acme GmbH", "Country/Region Key": "DE",
                 "VAT Registration No.": "DE136695976"}]
        body = (await client.post(URL + "?lookup_missing=grounded", json={"rows": rows})).json()
        assert body["rows"][0]["VAT_Found_Status"] is None
        assert body["summary"]["lookup"]["scope"] == "missing"

    @pytest.mark.asyncio
    async def test_file_experiment_compares_with_sap(self, client, lookups):
        content = _book([["Customer", "Name", "Country/Region Key", "VAT Registration No."],
                         ["1", "Acme GmbH", "DE", "DE136695976"], ["2", "Beta GmbH", "DE", None]])
        resp = await client.post(FILE_URL.replace("/file", "/llm-only-experiment/file"),
                                 files={"file": ("export.xlsx", content, XLSX)})
        wb = load_workbook(io.BytesIO(resp.content))
        _, out = _sheet(wb, "Data")
        assert [r["VAT_Found_vs_SAP"] for r in out] == ["MATCHES_SAP", "SAP_BLANK"]
        text = [str(c) for r in wb["Summary"].iter_rows(values_only=True) for c in r if c]
        assert any("every DE row" in c for c in text)
        assert any(c.startswith("VAT_Found_vs_SAP") for c in text)

    @pytest.mark.parametrize("proposal, sap, expected", [
        ("DE 136 695 976", "DE136695976", "MATCHES_SAP"),
        ("DE811128135", "DE136695976", "DIFFERS_FROM_SAP"),
        (None, "DE136695976", "NO_PROPOSAL"),
        ("DE136695976", None, "SAP_BLANK"),
    ])
    def test_vs_sap(self, proposal, sap, expected):
        from vat.lookup import vs_sap
        assert vs_sap(proposal, sap, "DE") == expected
