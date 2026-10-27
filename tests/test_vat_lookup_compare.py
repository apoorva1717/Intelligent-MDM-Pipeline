"""tools/vat_lookup_compare.py — outcome classification, totals and sheets.

The two lookup engines are fakes that propose a fixed value per Customer, so
every outcome class is exercised deterministically; VIES (the audit check) is
the scripted stub. The fixture is tests/fixtures/vat/VAT_TEST_EXPORT_30.xlsx.
"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
import pytest_asyncio
from openpyxl import load_workbook

from config import get_settings
from tests.mocks import vies_mock
from tests.mocks.vies_mock import ViesStub
from tools import vat_lookup_compare as compare_tool
from utils.cache import EvidenceCache
from vat.lookup import FoundResult
from vat.vies_client import VIESClient

FIXTURE = Path(__file__).parent / "fixtures" / "vat" / "VAT_TEST_EXPORT_30.xlsx"

# Rows 2–7: Google DE (2), Siemens (3), Bosch x2 (4, 5), Zalando (6), BASF (7).
A_ANSWERS = {
    "99000001": FoundResult("grounded", "FOUND_VERIFIED", value="DE811569869", tier="PROPOSED",
                            flags=("ATTRIBUTION_MATCH", "CANDIDATE:DE811569869|Google Germany GmbH|google.de|FOUND_VERIFIED")),
    "99000002": FoundResult("grounded", "FOUND_UNVERIFIED", value="DE129274235", tier="LIKELY",  # fake
                            flags=("ATTRIBUTION_PARTIAL", "CANDIDATE:a", "CANDIDATE:b")),
    "99000003": FoundResult("grounded", "FOUND_VERIFIED", value="DE129274235", tier="REVIEW",  # wrong
                            flags=("ATTRIBUTION_MATCH",)),
    "99000004": FoundResult("grounded", "NOT_FOUND", flags=("NO_IMPRESSUM_PAGE",)),
    "99000005": FoundResult("grounded", "FOUND_UNVERIFIED", value="DE149145248", tier="REVIEW",  # bad checksum
                            flags=("FORMAT_ERROR",)),
}
B_ANSWERS = {
    "99000001": FoundResult("llm_only", "LLM_GUESS:FORMAT_OK:VALID", value="DE811569869", evidence="high"),
    "99000002": FoundResult("llm_only", "LLM_GUESS:FORMAT_OK:VALID", value="DE811128135", evidence="high"),
    "99000003": FoundResult("llm_only", "LLM_GUESS:FORMAT_ERROR:NOT_CHECKED", value="DE811128136", evidence="medium"),
    "99000004": FoundResult("llm_only", "LLM_GUESS:FORMAT_OK:INVALID", value="DE136695976", evidence="low"),
    "99000005": FoundResult("llm_only", "LLM_GUESS:MISSING:NOT_CHECKED", value=None, evidence=None),
}


class _Fake:
    def __init__(self, mode, answers):
        self.mode, self.answers, self.seen = mode, answers, []

    async def lookup(self, inp):
        self.seen.append(inp.customer)
        return self.answers.get(inp.customer) or FoundResult(self.mode, "NOT_FOUND")

    def stats(self):
        return {"lookups": len(self.seen)}

    async def aclose(self):
        pass


@pytest_asyncio.fixture
async def report(tmp_path):
    settings = replace(get_settings(), vies_backoff_seconds=0.0, vies_min_interval_ms=0,
                       vies_requester_vat="", cache_frozen=False)
    engines = {"grounded": _Fake("grounded", A_ANSWERS), "llm_only": _Fake("llm_only", B_ANSWERS)}

    def make_client():
        return VIESClient(settings, cache=EvidenceCache(tmp_path),
                          transport=ViesStub(vies_mock.echo_valid_undisclosed).transport)

    wb = load_workbook(FIXTURE)
    result = await compare_tool.compare(
        wb, settings, rows=set(range(2, 8)), make_client=make_client,
        make_engine=lambda mode, s, c: engines[mode],
    )
    return wb, result, engines


class TestClassification:
    @pytest.mark.asyncio
    async def test_outcomes_per_row(self, report):
        _, result, _ = report
        rows = {e["sheet_row"]: e for e in result["rows"]}
        a = {r: rows[r]["grounded"]["outcome"] for r in rows}
        b = {r: rows[r]["llm_only"]["outcome"] for r in rows}
        assert a == {2: "CORRECT", 3: "WRONG_VIES_VALID", 4: "WRONG_VIES_VALID", 5: "WRONG_VIES_VALID",
                     6: "NULL", 7: "WRONG_FORMAT", 20: "NO_TRUTH"}
        assert b == {2: "CORRECT", 3: "WRONG_VIES_VALID", 4: "WRONG_FORMAT", 5: "WRONG_FORMAT",
                     6: "WRONG_PLAUSIBLE", 7: "NULL", 20: "NO_TRUTH"}

    @pytest.mark.asyncio
    async def test_mode_a_attribution_and_candidates(self, report):
        _, result, _ = report
        row3 = next(e for e in result["rows"] if e["sheet_row"] == 3)["grounded"]
        assert (row3["attribution"], row3["candidates"]) == ("ATTRIBUTION_PARTIAL", 2)

    @pytest.mark.asyncio
    async def test_gate_failed_wrong_values_get_an_audit_vies_check(self, report):
        _, result, _ = report
        row3 = next(e for e in result["rows"] if e["sheet_row"] == 3)["grounded"]
        assert row3["vies"] == "VALID"  # Mode A never asked; the audit did
        assert result["costs"]["audit"]["vies_checks"] >= 1


class TestTotals:
    @pytest.mark.asyncio
    async def test_each_customer_once(self, report):
        _, result, engines = report
        assert engines["grounded"].seen.count("99000003") == 1  # Bosch: 2 rows, 1 lookup
        a, b = result["totals"]["grounded"], result["totals"]["llm_only"]
        assert a["customers"] == b["customers"] == 5
        assert (a["CORRECT"], a["WRONG_VIES_VALID"], a["WRONG_FORMAT"], a["NULL"]) == (1, 2, 1, 1)
        assert (a["verified"], a["verified_but_wrong"]) == (2, 1)
        assert (b["CORRECT"], b["WRONG_VIES_VALID"], b["WRONG_PLAUSIBLE"], b["WRONG_FORMAT"], b["NULL"]) == (
            1, 1, 1, 1, 1)
        assert b["wrong_passes_format"] == 2

    @pytest.mark.asyncio
    async def test_correctness_by_confidence(self, report):
        _, result, _ = report
        c = result["confidence"]
        assert (c["high"]["customers"], c["high"]["CORRECT"], c["high"]["WRONG_VIES_VALID"]) == (2, 1, 1)
        assert (c["medium"]["WRONG_FORMAT"], c["low"]["WRONG_PLAUSIBLE"], c["null"]["NULL"]) == (1, 1, 1)


class TestWorkbook:
    @pytest.mark.asyncio
    async def test_sheets_and_untouched_data(self, report):
        wb, _, _ = report
        assert {"Comparison", "Cost", "Mode A rows", "Mode B rows"} <= set(wb.sheetnames)
        original = load_workbook(FIXTURE)["Data"]
        assert [r for r in wb["Data"].iter_rows(values_only=True)] == [
            r for r in original.iter_rows(values_only=True)]
        text = [c for r in wb["Comparison"].iter_rows(values_only=True) for c in r if c]
        assert any("EXPERIMENT" in str(c) for c in text)
        assert any("not an independent measure" in str(c) for c in text)
        header = next(r for r in wb["Comparison"].iter_rows(values_only=True) if r[0] == "Sheet row")
        for col in ("A attribution", "A candidates read", "B confidence", "B check digit ok", "B correct?"):
            assert col in header

    def test_parse_rows(self):
        assert compare_tool.parse_rows("2-4,9") == {2, 3, 4, 9}
        assert compare_tool.parse_rows(None) == set()


class TestProposals:
    @pytest.mark.parametrize("status,flags,expected", [
        ("FOUND_VERIFIED", "ATTRIBUTION_MATCH", True),
        ("FOUND_UNVERIFIED:VIES_SERVICE_UNAVAILABLE", "ATTRIBUTION_QUALIFIER; VIES:MS_UNAVAILABLE", True),
        ("FOUND_UNVERIFIED", "ATTRIBUTION_MISMATCH; CANDIDATE:x|y|z|FOUND_UNVERIFIED:ATTRIBUTION_MISMATCH", False),
        ("FOUND_UNVERIFIED", "ATTRIBUTION_PARTIAL", False),
        ("FOUND_UNVERIFIED", "DIRECTORY_DOMAIN; ATTRIBUTION_MATCH", False),  # a gate after attribution failed
        ("FOUND_VERIFIED", "CANDIDATE:x|y|z|FOUND_VERIFIED", False),  # no verdict of its own
        ("NOT_FOUND", "NO_IMPRESSUM_PAGE", False),
        ("LOOKUP_INCOMPLETE", "PAGE_UNAVAILABLE:403", False),
        (None, None, False),
    ])
    def test_is_proposal(self, status, flags, expected):
        assert compare_tool.is_proposal(status, flags) is expected

    @pytest.mark.asyncio
    async def test_totals_split_proposals_from_any_value(self, report):
        _, result, _ = report
        a = result["totals"]["grounded"]
        # 99000001 verified MATCH, right; 99000003 verified MATCH, wrong; the rest no proposal.
        assert (a["proposed"], a["proposed_correct"], a["proposed_wrong"], a["no_proposal"]) == (2, 1, 1, 3)
        assert (a["precision"], a["recall"]) == (0.5, 0.2)
        assert (a["any_value"], a["CORRECT"]) == (4, 1)  # old metric, kept

    @pytest.mark.asyncio
    async def test_sheet_has_proposal_block(self, report):
        wb, _, _ = report
        labels = [r[0] for r in wb["Comparison"].iter_rows(values_only=True) if r and r[0]]
        for label in ("Proposed", "No proposal", "Precision (Correct / Proposed)",
                      "Old metric: any value returned (gate passed or not)"):
            assert label in labels


class TestScoreExisting:
    @staticmethod
    def _workbook():
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Customer", "Name", "Country/Region Key", "City", "VAT Actual", "VAT Registration No.",
                   "VAT_Lookup_Mode", "VAT_Found_Value", "VAT_Found_Status", "VAT_Found_Flags",
                   "VAT_Found_Evidence", "VAT_Grounded_Value", "VAT_Grounded_Status", "VAT_Grounded_Flags"])
        ws.append(["1", "Google Germany GmbH", "DE", "HAMBURG", "DE 811569869", None, "LLM_ONLY",
                   "DE811569869", "LLM_GUESS:FORMAT_OK:VALID", None, "high",
                   "DE811569869", "FOUND_UNVERIFIED:VIES_SERVICE_UNAVAILABLE", "ATTRIBUTION_MATCH"])
        ws.append(["2", "ACL GmbH", "DE", "ROTTENBURG", "DE210202605", None, "LLM_ONLY",
                   None, "LLM_GUESS:MISSING:NOT_CHECKED", None, None,
                   "DE191570067", "FOUND_UNVERIFIED:VIES_SERVICE_UNAVAILABLE", "ATTRIBUTION_MATCH"])
        ws.append(["3", "Siemens AG", "DE", "MUENCHEN", "DE129274202", None, "LLM_ONLY",
                   None, "LLM_GUESS:MISSING:NOT_CHECKED", None, None,
                   "DE811569869", "FOUND_UNVERIFIED", "ATTRIBUTION_MISMATCH"])
        ws.append(["4", "Nobody GmbH", "DE", "BERLIN", None, None, "LLM_ONLY",
                   None, "LLM_GUESS:MISSING:NOT_CHECKED", None, None,
                   None, "LOOKUP_INCOMPLETE", "PAGE_UNAVAILABLE:403"])
        return wb

    @pytest.mark.asyncio
    async def test_scores_grounded_columns_without_lookups(self):
        wb = self._workbook()
        result = await compare_tool.score_results(wb)
        a, b = result["totals"]["grounded"], result["totals"]["llm_only"]
        assert a["customers"] == 3  # customer 4 has no SAP VAT ID
        assert (a["proposed"], a["proposed_correct"], a["proposed_wrong"], a["no_proposal"]) == (2, 1, 1, 1)
        assert (a["any_value"], a["CORRECT"]) == (3, 1)
        assert (b["CORRECT"], b["NULL"]) == (1, 2)
        assert result["costs"]["scored"]["vies_checks"] == 0
        assert {"Comparison", "Mode A rows", "Mode B rows"} <= set(wb.sheetnames)

    @pytest.mark.asyncio
    async def test_missing_truth_column_is_an_error(self):
        with pytest.raises(ValueError, match="No column"):
            await compare_tool.score_results(self._workbook(), truth_column="SAP VAT")


class TestTiers:
    @pytest.mark.asyncio
    async def test_per_tier_and_cumulative(self, report):
        _, result, _ = report
        tiers = result["totals"]["grounded"]["tiers"]
        # Customers with truth: 99000001 PROPOSED right, 99000002 LIKELY wrong,
        # 99000003 REVIEW wrong, 99000004 none, 99000005 REVIEW wrong.
        assert tiers["PROPOSED"] == {"n": 1, "correct": 1, "wrong": 0, "precision": 1.0, "recall": 0.2}
        assert tiers["LIKELY"] == {"n": 1, "correct": 0, "wrong": 1, "precision": 0.0, "recall": 0.0}
        assert (tiers["REVIEW"]["n"], tiers["REVIEW"]["correct"]) == (2, 0)
        assert tiers["PROPOSED + LIKELY"] == {"n": 2, "correct": 1, "wrong": 1, "precision": 0.5,
                                              "recall": 0.2}

    @pytest.mark.asyncio
    async def test_files_without_tiers_have_no_tier_block(self):
        result = await compare_tool.score_results(TestScoreExisting._workbook())
        assert result["totals"]["grounded"]["tiers"] is None

    @pytest.mark.asyncio
    async def test_score_existing_reads_the_grounded_tier(self):
        wb = TestScoreExisting._workbook()
        ws = wb["Data"]
        ws.cell(1, ws.max_column + 1, "VAT_Grounded_Tier")
        for row, tier in zip(range(2, 6), ("PROPOSED", "REVIEW", "REVIEW", None)):
            ws.cell(row, ws.max_column, tier)
        tiers = (await compare_tool.score_results(wb))["totals"]["grounded"]["tiers"]
        assert (tiers["PROPOSED"]["n"], tiers["PROPOSED"]["correct"], tiers["REVIEW"]["n"]) == (1, 1, 2)


class TestLlmWeb:
    @pytest.mark.asyncio
    async def test_side_by_side_file_scores_both_tiered_modes(self):
        wb = TestScoreExisting._workbook()
        ws = wb["Data"]
        header = [c.value for c in ws[1]]
        mode_col = header.index("VAT_Lookup_Mode") + 1
        for row in range(2, 6):
            ws.cell(row, mode_col, "LLM_WEB")
        ws.cell(1, ws.max_column + 1, "VAT_Found_Tier")
        found_tier = ws.max_column
        ws.cell(1, found_tier + 1, "VAT_Grounded_Tier")
        for row, (w, g) in zip(range(2, 6), (("PROPOSED", "PROPOSED"), (None, "REVIEW"),
                                            (None, "REVIEW"), (None, None))):
            ws.cell(row, found_tier, w)
            ws.cell(row, found_tier + 1, g)
        # Mode W's row 2 is a verified, attributed proposal of the right number.
        ws.cell(2, header.index("VAT_Found_Status") + 1, "FOUND_VERIFIED")
        ws.cell(2, header.index("VAT_Found_Flags") + 1, "ATTRIBUTION_MATCH; LLM_WEB_SEARCH")
        result = await compare_tool.score_results(wb)
        assert set(result["totals"]) == {"grounded", "llm_web"}
        w = result["totals"]["llm_web"]
        assert (w["tiers"]["PROPOSED"]["n"], w["tiers"]["PROPOSED"]["correct"]) == (1, 1)
        assert w["proposed"] == 1 and result["confidence"] == {}
        assert "Mode W rows" in wb.sheetnames
        header_row = next(r for r in wb["Comparison"].iter_rows(values_only=True) if r[0] == "Sheet row")
        assert "W tier" in header_row and "A tier" in header_row
