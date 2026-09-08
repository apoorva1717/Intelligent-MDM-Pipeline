"""Pass 18 — compute the evaluation results from ``data/eval/``.

Read-only. Nothing in the repository is written, no network call is made, no
LLM is invoked: every number this prints is arithmetic over the workbooks in
``data/eval/`` plus the recorded run artefacts under ``logs/``.

Usage::

    python tools/eval_report.py                 # every section
    python tools/eval_report.py --section issues

Why the issue counts are RECOMPUTED rather than read out of the workbooks
------------------------------------------------------------------------
Each stratum workbook carries an ``Issues`` column that some earlier run of
``POST /issues`` wrote (``api/routes.py:445`` appends it as the LAST column).
None of the ten records the commit it was produced at. The catalogue has moved
since: ``G1-NAME-001`` was withdrawn on 2026-09-07 and cannot be emitted at
this commit (``enrichment/issue_detection.py:270-274``). So the headline
figures here are produced by running the shipped detector over the shipped
cells — the same ``_parse_xlsx`` / ``_audit_rows`` path ``POST /issues`` uses
(``api/routes.py:227``, ``:762``) — and the shipped column is reported beside
them as a divergence check, per code.

The stress workbooks carry ``gt_dup_group`` / ``gt_expected_action``, which
``eval/dedup_eval.py`` cannot read (it wants ``expected_cluster`` /
``expected_routing``, ``eval/dedup_eval.py:45-46``). The adapter in
``_gt_routing`` states the mapping in one place and then hands the rows to the
repository's own scorer, so the pairwise numbers are that scorer's, not a
second implementation of it.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path
from typing import Any, Iterable

_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT))

from api.models import EnrichmentRecord  # noqa: E402
from api.output_columns import RESPONSE_COLUMNS  # noqa: E402
from api.routes import _audit_rows, _cell, _parse_xlsx  # noqa: E402
from enrichment.issue_detection import (  # noqa: E402
    EMITTED_CODES,
    ISSUE_CATALOGUE,
    QUALITY_GROUPS,
    VERIFICATION_GROUPS,
    issue_group,
)
from eval.dedup_eval import (  # noqa: E402
    business_risk_metrics,
    election_metrics,
    load_scored_rows,
    pairwise_metrics,
)
from utils.name_slots import RECORD_NAME_FIELDS  # noqa: E402
from utils.text_utils import is_blank  # noqa: E402

EVAL = _ROOT / "data" / "eval"
STRATA = ("S1", "S2", "S3", "S4", "S5")

# ── the field set the completeness KPI is measured over ──────────────────
# EnrichmentRecord's own Name block (api/models.py:90-110) and Address block
# (api/models.py:112-157). RECORD_NAME_FIELDS is the repository's name-slot
# tuple (utils/name_slots.py:43-46); the address list is the block's fields in
# declaration order. Nothing is invented here and nothing is dropped: the
# post-only sub-location slots are reported in their own table below because
# they exist on EnrichmentResult, not on EnrichmentRecord.
NAME_FIELDS: tuple[str, ...] = RECORD_NAME_FIELDS
ADDRESS_FIELDS: tuple[str, ...] = (
    "street_1", "house_number", "street_2", "street_3", "street_4",
    "street_5", "po_box", "country_region_key", "postal_code", "city",
    "region",
)
KPI_FIELDS: tuple[str, ...] = NAME_FIELDS + ADDRESS_FIELDS

# Columns the enriched export adds that carry address content moved OUT of the
# Street block. Reported separately: they have no pre-side counterpart, so
# including them in a pre-vs-post fill rate would compare against zero.
POST_ONLY_SUBLOCATION = (
    "Suite", "Building", "Floor", "Room", "Unit", "Mail Stop",
    "Unloading Point", "Mail Code", "Care Of", "Contact", "Email",
    "Operating Name", "Suggested Name", "Domain", "Department Domain",
    "Record Type", "ROR ID", "LEI ID", "Search Term 1", "Search Term 2",
)

LIVE_CODES: tuple[str, ...] = tuple(EMITTED_CODES)
MANDATORY_CODES: tuple[str, ...] = tuple(
    c for c in LIVE_CODES if ISSUE_CATALOGUE[c].mandatory
)


def sap_label(field: str) -> str:
    """The SAP column header a field is read from — its first alias."""
    alias = EnrichmentRecord.model_fields[field].validation_alias
    choices = getattr(alias, "choices", None)
    return str(choices[0]) if choices else field


def segment(code: str) -> str:
    """The three report blocks of ``POST /issues/compare``.

    Reproduces ``segment()`` at ``api/routes.py:540-561`` exactly: decided from
    ``raised`` and ``remedy``, never from the group.
    """
    entry = ISSUE_CATALOGUE[code]
    if entry.raised == "enriched":
        return "Verification"
    if entry.remedy == "steward":
        return "Expected to persist"
    return "Reduced"


REDUCTION_SET: tuple[str, ...] = tuple(
    c for c in LIVE_CODES if segment(c) == "Reduced"
)


# ── loading ──────────────────────────────────────────────────────────────

class Audit:
    """One workbook, parsed and audited on the ``POST /issues`` path."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self.headers, self.row_dicts = _parse_xlsx(path.read_bytes())
        self.records, self.issues = _audit_rows(self.headers, self.row_dicts)
        self.issue_columns = [
            i for i, h in enumerate(self.headers) if h == "Issues"
        ]
        self.shipped = {i: self._shipped(i) for i in self.issue_columns}

    def _shipped(self, index: int) -> list[set[str]]:
        """One ``Issues`` column, by position, as per-row code sets.

        Read off the sheet by column index rather than through
        ``row_dicts``: ``_parse_xlsx`` keys rows by header and keeps the last
        NON-EMPTY cell for a repeated header (``api/routes.py:272-278``), so
        in a workbook carrying two ``Issues`` columns the dict view is a
        cell-by-cell blend of both. No detector reads the column, so the blend
        does not affect detection — but it would silently corrupt this
        comparison.
        """
        from openpyxl import load_workbook
        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        out: list[set[str]] = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if all(c is None or str(c).strip() == "" for c in raw):
                continue
            cell = raw[index] if index < len(raw) else None
            out.append({c.strip() for c in str(cell or "").split(";")
                        if c.strip()})
        wb.close()
        return out

    @property
    def rows(self) -> int:
        return len(self.row_dicts)

    def code_counts(self) -> Counter:
        c: Counter = Counter()
        for codes in self.issues:
            c.update(codes)
        return c

    def group_counts(self) -> Counter:
        c: Counter = Counter()
        for codes in self.issues:
            c.update(issue_group(code) for code in codes)
        return c

    def records_with_issues(self) -> int:
        return sum(1 for codes in self.issues if codes)


def pairs() -> dict[str, tuple[Path, Path]]:
    """Resolve the five stratum pairs. Missing either side is fatal."""
    out: dict[str, tuple[Path, Path]] = {}
    missing: list[str] = []
    for s in STRATA:
        pre, post = EVAL / f"{s}_pre.xlsx", EVAL / f"{s}_post.xlsx"
        for p in (pre, post):
            if not p.exists():
                missing.append(str(p.relative_to(_ROOT)))
        out[s] = (pre, post)
    if missing:
        raise SystemExit(
            "STOP — stratum file(s) missing, no figure can be computed:\n  "
            + "\n  ".join(missing)
        )
    return out


# ── printing helpers ─────────────────────────────────────────────────────

def rule(title: str) -> None:
    print()
    print("=" * 78)
    print(title)
    print("=" * 78)


def table(headers: list[str], rows: list[list[Any]]) -> None:
    cells = [[str(c) for c in r] for r in rows]
    widths = [
        max([len(h)] + [len(r[i]) for r in cells]) for i, h in enumerate(headers)
    ]
    print("  " + "  ".join(h.ljust(widths[i]) for i, h in enumerate(headers)))
    print("  " + "  ".join("-" * w for w in widths))
    for r in cells:
        print("  " + "  ".join(r[i].ljust(widths[i]) for i in range(len(headers))))


def pct(after: int, before: int) -> str:
    if before == 0:
        return "n/a"
    return f"{(before - after) / before * 100:+.1f}%".replace("+", "-", 1) \
        if after <= before else f"+{(after - before) / before * 100:.1f}%"


def rel(before: int, after: int) -> str:
    """Relative reduction, positive = fewer issues after."""
    if before == 0:
        return "n/a"
    return f"{(before - after) / before * 100:.1f}%"


# ── section 1: pairing ───────────────────────────────────────────────────

def section_pairing(audits: dict[str, dict[str, Audit]]) -> None:
    rule("1 — FILE PAIRING AND ROW COUNTS")
    rows = []
    for s in STRATA:
        a, b = audits[s]["pre"], audits[s]["post"]
        rows.append([
            s,
            a.path.relative_to(_ROOT), a.rows, len(a.headers), len(a.records),
            b.path.relative_to(_ROOT), b.rows, len(b.headers), len(b.records),
        ])
    table(
        ["stratum", "pre file", "rows", "cols", "records",
         "post file", "rows", "cols", "records"],
        rows,
    )
    print()
    print("  Every stratum resolves to both files; no stratum is skipped.")


# ── section 2: vocabulary ────────────────────────────────────────────────

def section_vocabulary() -> None:
    rule("2 — VOCABULARY")
    status = Counter(e.status for e in ISSUE_CATALOGUE.values())
    print(f"  ISSUE_CATALOGUE declared : {len(ISSUE_CATALOGUE)}")
    print(f"  live / emittable         : {len(LIVE_CODES)}  {dict(status)}")
    print(f"  quality groups           : {list(QUALITY_GROUPS)}")
    print(f"  verification groups      : {list(VERIFICATION_GROUPS)}"
          "  (never counted in a reduction figure)")
    print()
    print(f"  Reduction set (segment == 'Reduced'): {len(REDUCTION_SET)} codes")
    by_group: dict[str, list[str]] = defaultdict(list)
    for c in REDUCTION_SET:
        by_group[issue_group(c)].append(c)
    for g in sorted(by_group):
        print(f"    {g}  ({len(by_group[g])})  {', '.join(by_group[g])}")
    print()
    print(f"  Mandatory (DATAshaper Error) live codes: {len(MANDATORY_CODES)}")
    for c in MANDATORY_CODES:
        print(f"    {c}  {ISSUE_CATALOGUE[c].name}  [{segment(c)}]")
    print()
    print("  DedupIssue codes from /api/dedup/score are cluster-quality")
    print("  diagnostics and are reported in section 7 only. They are never")
    print("  summed with catalogue codes.")


# ── section 3: issues, pre vs post ───────────────────────────────────────

def section_issues(audits: dict[str, dict[str, Audit]]) -> None:
    rule("3 — ISSUES PER STRATUM, PRE VS POST (recomputed at this commit)")

    totals_pre: Counter = Counter()
    totals_post: Counter = Counter()

    for s in STRATA:
        pre, post = audits[s]["pre"], audits[s]["post"]
        cp, cq = pre.code_counts(), post.code_counts()
        totals_pre.update(cp)
        totals_post.update(cq)

        print()
        print(f"--- {s} " + "-" * (74 - len(s)))
        print(f"  rows: pre {pre.rows}  post {post.rows}")
        print(f"  records carrying >=1 issue: pre {pre.records_with_issues()}"
              f"  post {post.records_with_issues()}")

        # per code
        codes = sorted(set(cp) | set(cq))
        rows = [
            [c, issue_group(c), segment(c),
             "yes" if ISSUE_CATALOGUE[c].mandatory else "",
             cp.get(c, 0), cq.get(c, 0), cq.get(c, 0) - cp.get(c, 0),
             rel(cp.get(c, 0), cq.get(c, 0))]
            for c in codes
        ]
        print()
        table(["code", "grp", "segment", "mand", "pre", "post", "delta",
               "reduction"], rows)

        # per group
        gp, gq = pre.group_counts(), post.group_counts()
        groups = sorted(set(gp) | set(gq))
        print()
        table(
            ["group", "pre", "post", "delta", "reduction"],
            [[g, gp.get(g, 0), gq.get(g, 0), gq.get(g, 0) - gp.get(g, 0),
              rel(gp.get(g, 0), gq.get(g, 0))] for g in groups],
        )

        # segments + the two named sets
        print()
        seg_rows = []
        for name in ("Reduced", "Expected to persist", "Verification"):
            b = sum(v for k, v in cp.items() if segment(k) == name)
            a = sum(v for k, v in cq.items() if segment(k) == name)
            seg_rows.append([name, b, a, a - b, rel(b, a)])
        b = sum(cp.get(c, 0) for c in REDUCTION_SET)
        a = sum(cq.get(c, 0) for c in REDUCTION_SET)
        seg_rows.append([f"reduction set ({len(REDUCTION_SET)} codes)",
                         b, a, a - b, rel(b, a)])
        b = sum(cp.get(c, 0) for c in MANDATORY_CODES)
        a = sum(cq.get(c, 0) for c in MANDATORY_CODES)
        seg_rows.append([f"mandatory ({len(MANDATORY_CODES)} codes)",
                         b, a, a - b, rel(b, a)])
        b, a = sum(cp.values()), sum(cq.values())
        seg_rows.append(["all live codes", b, a, a - b, rel(b, a)])
        table(["set", "pre", "post", "delta", "reduction"], seg_rows)

    # column gating — why a mandatory count can grow without a regression
    print()
    print("--- column gating on the required-field rules " + "-" * 32)
    print("  The five G2-VAL-* rules fire only when the column EXISTS in the")
    print("  file and is blank (enrichment/issue_detection.py:1189-1215).")
    print("  A code whose column is absent from the pre file cannot fire there,")
    print("  so its pre count is 0 by construction, not by cleanliness.")
    print()
    from api.routes import _present_fields
    req = {
        "G2-VAL-001": "name_1", "G2-VAL-002": "postal_code",
        "G2-VAL-004": "region", "G2-VAL-007": "search_term_1",
        "G2-VAL-008": "country_region_key",
    }
    rows = []
    for code, field in req.items():
        row = [code, sap_label(field)]
        for s in STRATA:
            pre_f = _present_fields(audits[s]["pre"].headers)
            post_f = _present_fields(audits[s]["post"].headers)
            row.append(("y" if field in pre_f else "n") + "/"
                       + ("y" if field in post_f else "n"))
        rows.append(row)
    table(["code", "column"] + [f"{s} pre/post" for s in STRATA], rows)

    # all strata
    print()
    print("--- S1-S5 combined " + "-" * 59)
    codes = sorted(set(totals_pre) | set(totals_post))
    print()
    table(
        ["code", "grp", "segment", "pre", "post", "delta", "reduction"],
        [[c, issue_group(c), segment(c), totals_pre.get(c, 0),
          totals_post.get(c, 0), totals_post.get(c, 0) - totals_pre.get(c, 0),
          rel(totals_pre.get(c, 0), totals_post.get(c, 0))] for c in codes],
    )
    print()
    seg_rows = []
    for name in ("Reduced", "Expected to persist", "Verification"):
        b = sum(v for k, v in totals_pre.items() if segment(k) == name)
        a = sum(v for k, v in totals_post.items() if segment(k) == name)
        seg_rows.append([name, b, a, a - b, rel(b, a)])
    b = sum(totals_pre.get(c, 0) for c in REDUCTION_SET)
    a = sum(totals_post.get(c, 0) for c in REDUCTION_SET)
    seg_rows.append([f"reduction set ({len(REDUCTION_SET)})", b, a, a - b, rel(b, a)])
    b = sum(totals_pre.get(c, 0) for c in MANDATORY_CODES)
    a = sum(totals_post.get(c, 0) for c in MANDATORY_CODES)
    seg_rows.append([f"mandatory ({len(MANDATORY_CODES)})", b, a, a - b, rel(b, a)])
    seg_rows.append(["all live codes", sum(totals_pre.values()),
                     sum(totals_post.values()),
                     sum(totals_post.values()) - sum(totals_pre.values()),
                     rel(sum(totals_pre.values()), sum(totals_post.values()))])
    table(["set", "pre", "post", "delta", "reduction"], seg_rows)


# ── section 4: shipped column vs recomputation ───────────────────────────

def section_divergence(audits: dict[str, dict[str, Audit]]) -> None:
    rule("4 — SHIPPED `Issues` COLUMN VS THE DETECTOR AT THIS COMMIT")
    print("  Every `Issues` column in every workbook, compared as code SETS")
    print("  against the detector re-run over the same cells at this commit.")
    print("  `POST /issues` appends its column LAST (api/routes.py:445), so an")
    print("  earlier column of that name is input-side, not output.")
    print()
    rows = []
    for s in STRATA:
        for kind in ("pre", "post"):
            a = audits[s][kind]
            if not a.issue_columns:
                rows.append([f"{s}_{kind}", "-", "-", "-", "-",
                             "no Issues column"])
                continue
            for i in a.issue_columns:
                shipped = a.shipped[i]
                agree = sum(1 for sh, lv in zip(shipped, a.issues)
                            if sh == set(lv))
                only_ship: Counter = Counter()
                only_live: Counter = Counter()
                for sh, lv in zip(shipped, a.issues):
                    only_ship.update(sh - set(lv))
                    only_live.update(set(lv) - sh)
                rows.append([
                    f"{s}_{kind}", i,
                    "last" if i == a.issue_columns[-1] else "earlier",
                    f"{agree}/{a.rows}",
                    dict(only_ship.most_common()) or "-",
                    dict(only_live.most_common()) or "-",
                ])
    table(["workbook", "col (0-based)", "position", "rows agreeing",
           "shipped only", "detector only"], rows)


# ── section 5: completeness KPI ──────────────────────────────────────────

def section_completeness(audits: dict[str, dict[str, Audit]]) -> None:
    rule("5 — COMPLETENESS KPI (filled / expected, name + address block)")
    print("  Field set: EnrichmentRecord's Name block (api/models.py:90-110,")
    print("  utils/name_slots.py:43-46) and Address block (api/models.py:112-157).")
    print("  'expected' is the row count. 'col' says whether the workbook")
    print("  carries a column that maps onto the field at all: a field with no")
    print("  column can only read as 0 filled, which is a different fact from")
    print("  an empty cell.")

    for s in STRATA:
        pre, post = audits[s]["pre"], audits[s]["post"]
        pre_cols = _mapped_headers(pre)
        post_cols = _mapped_headers(post)
        print()
        print(f"--- {s} " + "-" * (74 - len(s)))
        rows = []
        for f in KPI_FIELDS:
            fp = sum(1 for r in pre.records if not is_blank(getattr(r, f)))
            fq = sum(1 for r in post.records if not is_blank(getattr(r, f)))
            rows.append([
                sap_label(f), f,
                "y" if f in pre_cols else "n", f"{fp}/{pre.rows}",
                "y" if f in post_cols else "n", f"{fq}/{post.rows}",
                fq - fp,
            ])
        table(["SAP column", "field", "col", "pre filled", "col",
               "post filled", "delta"], rows)

        both = [f for f in KPI_FIELDS if f in pre_cols and f in post_cols]
        fp = sum(1 for f in both for r in pre.records
                 if not is_blank(getattr(r, f)))
        fq = sum(1 for f in both for r in post.records
                 if not is_blank(getattr(r, f)))
        exp = len(both) * pre.rows
        print()
        print(f"  headline over the {len(both)} fields carried by BOTH files:")
        print(f"    pre  {fp}/{exp} = {fp / exp * 100:.1f}%")
        print(f"    post {fq}/{exp} = {fq / exp * 100:.1f}%")
        print(f"    delta {fq - fp:+d} filled cells "
              f"({(fq - fp) / exp * 100:+.1f} pp)")
        missing = [sap_label(f) for f in KPI_FIELDS if f not in pre_cols]
        if missing:
            print(f"    fields with no pre-side column: {', '.join(missing)}")

    # values present before and gone after, joined on the record id
    print()
    print("--- values lost: filled in pre, blank in post, joined on Customer "
          + "-" * 12)
    rows = []
    for s in STRATA:
        pre, post = audits[s]["pre"], audits[s]["post"]
        pre_by = {r.record_id: r for r in pre.records}
        post_by = {r.record_id: r for r in post.records}
        shared = [k for k in pre_by if k in post_by]
        row = [s, len(shared)]
        for f in KPI_FIELDS:
            row.append(sum(
                1 for k in shared
                if not is_blank(getattr(pre_by[k], f))
                and is_blank(getattr(post_by[k], f))
            ))
        rows.append(row)
    table(["stratum", "joined"] + [sap_label(f) for f in KPI_FIELDS], rows)
    print("  A count here is a value the pre file carried and the post file")
    print("  does not. For the Street block most of it is relocation into the")
    print("  sub-location columns below; for Name 1 there is no such")
    print("  destination in the schema.")

    # post-only columns
    print()
    print("--- columns the enriched export adds (no pre-side counterpart) "
          + "-" * 15)
    rows = []
    for s in STRATA:
        post = audits[s]["post"]
        idx = {h: i for i, h in enumerate(post.headers)}
        vals = []
        for col in POST_ONLY_SUBLOCATION:
            if col not in idx:
                vals.append("-")
                continue
            n = sum(1 for r in post.row_dicts if not is_blank(r.get(col)))
            vals.append(str(n))
        rows.append([s] + vals)
    table(["stratum"] + list(POST_ONLY_SUBLOCATION), rows)


def _mapped_headers(a: Audit) -> set[str]:
    from api.routes import _present_fields
    return _present_fields(a.headers)


# ── section 6: improvement index ─────────────────────────────────────────

def section_improvement(audits: dict[str, dict[str, Audit]]) -> None:
    rule("6 — IMPROVEMENT INDEX (post-run count / mean count across categories)")
    print("  'category' is read two ways because the specification does not")
    print("  fix one; both are given. An index above 1.0 is the next-iteration")
    print("  trigger.")

    # (a) issue group within a stratum
    for s in STRATA:
        post = audits[s]["post"]
        g = post.group_counts()
        cats = sorted(set(g) | set(QUALITY_GROUPS) | set(VERIFICATION_GROUPS))
        counts = [g.get(c, 0) for c in cats]
        mean = sum(counts) / len(cats) if cats else 0.0
        print()
        print(f"--- {s}: by issue group " + "-" * (57 - len(s)))
        rows = [[c, g.get(c, 0), f"{(g.get(c, 0) / mean if mean else 0):.2f}",
                 "ABOVE" if mean and g.get(c, 0) / mean > 1.0 else ""]
                for c in cats]
        table(["group", "post count", "index", ""], rows)
        print(f"  mean across {len(cats)} groups = {mean:.2f}"
              f"   above 1.0: "
              f"{', '.join(c for c in cats if mean and g.get(c, 0) / mean > 1.0) or 'none'}")

    # (b) stratum as the category, one table
    print()
    print("--- across strata (the stratum IS the record category) " + "-" * 23)
    counts = {s: sum(audits[s]['post'].code_counts().values()) for s in STRATA}
    mean = sum(counts.values()) / len(STRATA)
    cat = {s: audits[s]["post"].row_dicts[0].get("category", "?") for s in STRATA}
    table(
        ["stratum", "category", "post count", "index", ""],
        [[s, cat[s], counts[s], f"{counts[s] / mean:.2f}",
          "ABOVE" if counts[s] / mean > 1.0 else ""] for s in STRATA],
    )
    print(f"  mean across {len(STRATA)} strata = {mean:.2f}   above 1.0: "
          f"{', '.join(s for s in STRATA if counts[s] / mean > 1.0) or 'none'}")


# ── section 7: clustering ────────────────────────────────────────────────

_ACTION_TO_ROUTING = {
    # gt_expected_action (Method sheet, verbatim) -> dedup_eval's vocabulary.
    # Only MERGE asks for a collapse, so only MERGE forms ground-truth pairs;
    # REVIEW is the steward route, everything else must NOT collapse.
    "MERGE": "cluster",
    "REVIEW": "manual_review",
    "LINK": "unique",
    "DO NOT MERGE": "unique",
    "UNIQUE": "unique",
}


def _gt_routing(action: str | None) -> str | None:
    if not action:
        return None
    head = str(action).split(" - ")[0].split(":")[0].strip().upper()
    for key, value in _ACTION_TO_ROUTING.items():
        if head.startswith(key):
            return value
    return None


def _gt_table(path: Path, sheet: str) -> dict[str, dict[str, str]]:
    """{Customer: {gt_* column: value}} from a stress workbook sheet."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    want = [h for h in hdr if h.startswith("gt_")]
    out: dict[str, dict[str, str]] = {}
    for raw in rows[1:]:
        rec = dict(zip(hdr, raw))
        cid = rec.get("Customer")
        if cid is None:
            continue
        # duplicate gt_expected_action columns: first occurrence wins, which is
        # the column stress_200_scored.xlsx kept (its Reconciliation sheet).
        keep: dict[str, str] = {}
        for h, v in zip(hdr, raw):
            if h in want and h not in keep and v is not None:
                keep[h] = str(v).strip()
        out[str(cid).strip()] = keep
    wb.close()
    return out


def _attach_gt(rows: list[dict], gt: dict[str, dict[str, str]]) -> int:
    matched = 0
    for r in rows:
        g = gt.get(r["row_id"])
        if not g:
            r["expected_cluster"] = None
            r["expected_routing"] = None
            continue
        matched += 1
        r["expected_cluster"] = g.get("gt_dup_group")
        r["expected_routing"] = _gt_routing(g.get("gt_expected_action"))
        r["gt_entity_id"] = g.get("gt_entity_id")
        r["gt_trap_group"] = g.get("gt_trap_group")
        r["gt_difficulty"] = g.get("gt_difficulty")
    return matched


def _merge_breakdown(rows: list[dict]) -> dict[str, Any]:
    """Over- and under-merges, classified by what the truth actually said."""
    by_id = {r["row_id"]: r for r in rows}
    pred: dict[str, list[str]] = defaultdict(list)
    for r in rows:
        if r.get("cluster_id"):
            pred[r["cluster_id"]].append(r["row_id"])

    over: Counter = Counter()
    over_examples: dict[str, list[str]] = defaultdict(list)
    for members in pred.values():
        for a, b in combinations(sorted(members), 2):
            ra, rb = by_id[a], by_id[b]
            same_dup = (
                ra.get("expected_cluster")
                and ra["expected_cluster"] == rb.get("expected_cluster")
            )
            if same_dup:
                continue
            same_entity = (
                ra.get("gt_entity_id")
                and ra["gt_entity_id"] == rb.get("gt_entity_id")
            )
            same_trap = (
                ra.get("gt_trap_group")
                and ra["gt_trap_group"] == rb.get("gt_trap_group")
            )
            kind = ("link (same entity, different site)" if same_entity
                    else "TRAP (different organisation)" if same_trap
                    else "unrelated")
            over[kind] += 1
            over_examples[kind].append(f"{a}+{b}")

    truth: dict[str, list[str]] = defaultdict(list)
    for r in rows:
        if r.get("expected_routing") == "cluster" and r.get("expected_cluster"):
            truth[r["expected_cluster"]].append(r["row_id"])
    under: list[str] = []
    for members in truth.values():
        for a, b in combinations(sorted(members), 2):
            ca, cb = by_id[a].get("cluster_id"), by_id[b].get("cluster_id")
            if ca is None or cb is None or ca != cb:
                under.append(f"{a}+{b}")

    return {
        "over_merge_pairs": dict(over),
        "over_merge_examples": {k: v[:6] for k, v in over_examples.items()},
        "under_merge_pairs": len(under),
        "under_merge_examples": under[:12],
        "gt_merge_groups": sum(1 for m in truth.values() if len(m) >= 2),
        "predicted_clusters": len(pred),
    }


def _score_workbook(path: Path, gt_source: Path, gt_sheet: str,
                    label: str) -> None:
    print()
    print(f"--- {label} " + "-" * max(0, 74 - len(label)))
    print(f"  scored workbook : {path.relative_to(_ROOT)}")
    print(f"  ground truth    : {gt_source.relative_to(_ROOT)} [{gt_sheet}]")
    rows = load_scored_rows(path)
    gt = _gt_table(gt_source, gt_sheet)
    matched = _attach_gt(rows, gt)
    print(f"  rows scored     : {len(rows)}   ground truth matched: {matched}")
    unmatched = [r["row_id"] for r in rows if not gt.get(r["row_id"])]
    if unmatched:
        print(f"  rows with NO ground truth ({len(unmatched)}): "
              f"{', '.join(unmatched[:10])}"
              + (" ..." if len(unmatched) > 10 else ""))

    routing = Counter(r.get("routing") or "blank" for r in rows)
    print(f"  routing         : {dict(routing.most_common())}")
    exp_routing = Counter(r.get("expected_routing") or "unmapped" for r in rows)
    print(f"  expected routing: {dict(exp_routing.most_common())}")

    pw = pairwise_metrics(rows)
    print()
    table(["metric", "value"], [[k, v] for k, v in pw.items()])

    br = _merge_breakdown(rows)
    print()
    print(f"  ground-truth MERGE groups of size >= 2 : {br['gt_merge_groups']}")
    print(f"  clusters the run produced              : {br['predicted_clusters']}")
    print(f"  under-merge pairs (truth MERGE, not clustered): "
          f"{br['under_merge_pairs']}")
    if br["under_merge_examples"]:
        print(f"    e.g. {', '.join(br['under_merge_examples'])}")
    print(f"  over-merge pairs (clustered, truth not MERGE):")
    if br["over_merge_pairs"]:
        for k, v in br["over_merge_pairs"].items():
            print(f"    {k}: {v}   e.g. "
                  f"{', '.join(br['over_merge_examples'][k])}")
    else:
        print("    none")

    if any(r.get("election_status") for r in rows):
        print()
        em = election_metrics(rows)
        table(["election metric", "value"],
              [[k, v] for k, v in em.items() if k != "tiebreak_decided_clusters"]
              + [["tiebreak_decided_clusters",
                  em["tiebreak_decided_clusters"]["count"]]])
        print(f"    tie-break cluster ids: "
              f"{em['tiebreak_decided_clusters']['cluster_ids']}")
        print()
        brm = business_risk_metrics(rows)
        table(["business risk", "count", "row ids"],
              [[k, v["count"], ", ".join(v["row_ids"][:8]) or "-"]
               for k, v in brm.items()])
    else:
        print()
        print("  No election columns in this workbook — election and")
        print("  business-risk metrics are not computable from it.")


def section_clustering(audits: dict[str, dict[str, Audit]]) -> None:
    rule("7 — CLUSTERING AND ELECTION")
    print("  Ground-truth mapping onto eval/dedup_eval.py's vocabulary")
    print("  (the Method sheet's 'How to score' rule, stated once):")
    for k, v in _ACTION_TO_ROUTING.items():
        print(f"    gt_expected_action {k:<13} -> expected_routing {v}")
    print("    expected_cluster = gt_dup_group")
    print("  Only MERGE forms ground-truth pairs, so LINK and the two DO NOT")
    print("  MERGE actions can only ever appear as a false positive.")

    _score_workbook(
        EVAL / "dedup_STRESS_200_v1_enriched_dedup.xlsx",
        EVAL / "dedup_STRESS_200_v1_enriched_dedup.xlsx", "Sheet",
        "stress set, 200-row grain (clustering run, no scoring)",
    )
    _score_workbook(
        EVAL / "stress_200_scored.xlsx",
        EVAL / "dedup_STRESS_200_v1-verified.xlsx", "Data",
        "stress set, 183-row grain (clustering + scoring + election)",
    )

    print()
    print("--- S5 test set " + "-" * 62)
    post = audits["S5"]["post"]
    produced = [h for h in post.headers if h in ("Cluster ID", "Routing",
                                                 "election_status")]
    ann = Counter(r.get("cluster_id") for r in post.row_dicts
                  if r.get("cluster_id"))
    print(f"  annotation columns present : cluster_id, cluster_role")
    print(f"  annotated clusters         : {len(ann)} over "
          f"{sum(ann.values())} records, sizes "
          f"{dict(Counter(ann.values()))}")
    print(f"  cluster_role distribution  : "
          f"{dict(Counter((r.get('cluster_role') or 'blank') for r in post.row_dicts).most_common())}")
    print(f"  clustering OUTPUT columns  : {produced or 'NONE'}")
    print("  ** MEASUREMENT REQUIRED ** — no run output exists for S5, so")
    print("  expected-vs-found, over-merges, under-merges, manual_review and")
    print("  unique counts cannot be computed for this stratum. Producing them")
    print("  needs POST /api/dedup/cluster over S5_post.xlsx.")
    print("  Note also that eval/dedup_eval.py maps the header `cluster_id`")
    print("  onto its OUTPUT field (eval/dedup_eval.py:47), so pointing it at")
    print("  a stratum workbook would read the annotation as the run's answer.")


# ── section 8: cost ──────────────────────────────────────────────────────

_COST_KEYS = (
    "total", "tier1_resolved", "lei_attempts", "tier2a_population_count",
    "tier2a_verification_count", "tier2b_count", "tier3_count",
    "page_reads_attempted", "wikidata_queried", "liveness_ror_queried",
    "domain_from_serp", "contact_lookup_attempted",
    "evidence_network_calls", "evidence_cache_hits", "evidence_cache_frozen",
    "processing_time_ms",
)


def section_cost() -> None:
    rule("8 — COST")
    report = _ROOT / "logs" / "runs" / "determinism_S1_327ee53.json"
    if report.exists():
        data = json.loads(report.read_text())
        s = data["summary_run1"]
        print(f"  Recorded run summary — {report.relative_to(_ROOT)}")
        print("  (the only per-run counter set in the repository that names a")
        print("   stratum; the name is the only evidence it is S1)")
        print()
        rows = []
        for k in _COST_KEYS:
            v = s.get(k, "-")
            per = (f"{v / s['total']:.2f}"
                   if isinstance(v, int) and not isinstance(v, bool)
                   and k != "total" else "")
            rows.append([k, v, per])
        table(["counter", "value", "per record"], rows)
        print()
        print("  evidence_network_calls = 0 and evidence_cache_frozen = true:")
        print("  this run replayed a recorded cache, so it measures the WORK")
        print("  the pipeline asked for, not calls billed on that run.")
        print(f"  network calls by namespace: "
              f"{s.get('evidence_network_calls_by_namespace')}")
    else:
        print(f"  ** MEASUREMENT REQUIRED ** — {report} absent.")

    print()
    print("  Dedup LLM calls — data/eval/stress_200_scored.xlsx :: Run sheet")
    from openpyxl import load_workbook
    wb = load_workbook(EVAL / "stress_200_scored.xlsx", read_only=True,
                       data_only=True)
    run = {str(r[0]): r[1] for r in wb["Run"].iter_rows(values_only=True)
           if r[0] is not None}
    wb.close()
    table(["setting", "value"],
          [[k, v] for k, v in run.items() if k != "setting"])
    if run.get("llm_calls") and run.get("rows_in"):
        print(f"  LLM calls per record: "
              f"{int(run['llm_calls']) / int(run['rows_in']):.3f}")

    print()
    print("  Recorded evidence caches (unique keys per namespace, NOT calls —")
    print("  a repeated question is one key). None is attributable to a")
    print("  stratum: no cache directory records the batch it was recorded on.")
    cache = _ROOT / "logs" / "cache"
    if cache.is_dir():
        namespaces = sorted({n.name for d in cache.iterdir() if d.is_dir()
                             for n in d.iterdir() if n.is_dir()})
        rows = []
        for d in sorted(cache.iterdir()):
            if not d.is_dir():
                continue
            rows.append([d.name] + [
                str(len(list((d / n).iterdir()))) if (d / n).is_dir() else "-"
                for n in namespaces
            ])
        table(["cache dir"] + namespaces, rows)

    print()
    print("  ** MEASUREMENT REQUIRED ** — cost in currency. No price, rate")
    print("  card or token-price constant exists anywhere in this repository")
    print("  (grep for price/pricing/USD/cost_per over *.py, *.json, *.md,")
    print("  *.yaml returns no rate). Supplying it needs the SerpAPI plan")
    print("  rate and the Azure OpenAI deployment's per-1K-token price for")
    print("  the deployment named on the Run sheet, or an App Insights cost")
    print("  export. Neither is in the repository.")
    print("  ** MEASUREMENT REQUIRED ** — calls per record per tier per")
    print("  stratum for S2-S5: no run summary exists for those strata.")


# ── section 9: determinism ───────────────────────────────────────────────

_RUN_DIFF_PAIRS = (
    # (label, run1, run2) — the tracked scripts/run_batch.py --json artefacts
    # under eval/out/, which are keyed by the commit that produced them
    # (eval/out/RUNS.md). Bare files are d3a3cfc.
    ("S1  d3a3cfc -> f57782f", "eval/out/S1_results.json",
     "eval/out/f57782f/S1_results.json"),
    ("S1  f57782f -> 327ee53", "eval/out/f57782f/S1_results.json",
     "eval/out/327ee53/S1_results.json"),
    ("S4  d3a3cfc -> f57782f", "eval/out/S4_results.json",
     "eval/out/f57782f/S4_results.json"),
    ("S4  f57782f -> 327ee53", "eval/out/f57782f/S4_results.json",
     "eval/out/327ee53/S4_results.json"),
)


def section_determinism() -> None:
    rule("9 — DETERMINISM AND RUN-TO-RUN STABILITY")
    from tools.run_diff import render

    print("  9a — tools/run_diff.py RUN HERE, at this commit, over the tracked")
    print("  run artefacts in eval/out/. Each pair is the SAME batch at two")
    print("  DIFFERENT commits, so a difference is a code change, not")
    print("  non-determinism: this measures how far the output moved between")
    print("  commits, not whether one commit repeats itself.")
    import subprocess
    for label, a, b in _RUN_DIFF_PAIRS:
        pa, pb = _ROOT / a, _ROOT / b
        print()
        print(f"--- {label} " + "-" * max(0, 74 - len(label)))
        if not (pa.exists() and pb.exists()):
            print(f"  ** MEASUREMENT REQUIRED ** — missing {a} or {b}")
            continue
        print(f"$ python tools/run_diff.py {a} {b} --quiet")
        proc = subprocess.run(
            [sys.executable, str(_ROOT / "tools" / "run_diff.py"),
             str(pa), str(pb), "--quiet"],
            capture_output=True, text=True, cwd=str(_ROOT),
        )
        print(proc.stdout.rstrip())
        for line in proc.stderr.splitlines():
            if "Warning" not in line and "warnings.warn" not in line:
                print(line)
        print(f"[exit {proc.returncode}]")

    print()
    print("  9b — the determinism gate proper: TWO RUNS AT ONE COMMIT. The two")
    print("  run artefacts of each such pair are not in the repository, only")
    print("  run_diff's stored --json report of them. Each is re-rendered")
    print("  verbatim below through tools.run_diff.render().")
    reports = sorted((_ROOT / "logs" / "runs").glob("determinism_S1*.json"))
    if not reports:
        print("  ** MEASUREMENT REQUIRED ** — no run_diff report artefact.")
        return
    for path in reports:
        data = json.loads(path.read_text())
        print()
        print(f"--- {path.relative_to(_ROOT)} " + "-" * 20)
        print(render(data, verbose=True))
        print(f"\nrun 2 network calls: "
              f"{data['summary_run2'].get('evidence_network_calls')}")


# ── section 10: cross-check against the recorded runs ────────────────────

_RECORDED_RUNS = {
    "S1": "eval/out/327ee53/S1_results.json",
    "S4": "eval/out/327ee53/S4_results.json",
    "S5": "eval/out/327ee53/S5_results.json",
}


def _audit_run_json(path: Path) -> tuple[list, list[list[str]], list[dict]]:
    """Audit a scripts/run_batch.py --json artefact as if it were the workbook.

    The rows are rebuilt through the shipped output schema and the shipped
    cell adapter — RESPONSE_COLUMNS (api/output_columns.py:22) and _cell
    (api/routes.py:372-382) — so the dict handed to the detector is the one
    _build_output_xlsx would have written to the sheet.
    """
    data = json.loads(path.read_text())
    headers = list(RESPONSE_COLUMNS.values())
    row_dicts: list[dict[str, str]] = []
    for result in data["results"]:
        row: dict[str, str] = {}
        for field, header in RESPONSE_COLUMNS.items():
            value = _cell(result.get(field))
            if value is None:
                continue
            text = str(value).strip()
            if text:
                row[header] = text
        row_dicts.append(row)
    records, issues = _audit_rows(headers, row_dicts)
    return records, issues, row_dicts


def section_recorded(audits: dict[str, dict[str, Audit]]) -> None:
    rule("10 — CROSS-CHECK AGAINST THE RECORDED RUNS IN eval/out/")
    print("  eval/out/327ee53/ holds the scripts/run_batch.py --json artefact")
    print("  of a frozen run of S1, S4 and S5 at commit 327ee53")
    print("  (eval/out/RUNS.md). Its record ids are the same 100 per stratum as")
    print("  the data/eval workbooks. Auditing that artefact gives a SECOND")
    print("  post-enrichment state for those three strata — one that names the")
    print("  commit it was produced at, which no post workbook does.")
    print("  A difference below is a difference between two artefacts produced")
    print("  at two commits; it is not by itself a regression at either.")

    for s, rel in _RECORDED_RUNS.items():
        path = _ROOT / rel
        print()
        print(f"--- {s} " + "-" * (74 - len(s)))
        if not path.exists():
            print(f"  ** MEASUREMENT REQUIRED ** — {rel} absent.")
            continue
        recs, issues, rows = _audit_run_json(path)
        post = audits[s]["post"]
        run_ids = {r.record_id for r in recs}
        wb_ids = {r.record_id for r in post.records}
        print(f"  {rel}")
        print(f"  records {len(recs)}   ids shared with {s}_post.xlsx: "
              f"{len(run_ids & wb_ids)}")

        run_counts: Counter = Counter()
        for codes in issues:
            run_counts.update(codes)
        wb_counts = post.code_counts()
        codes = sorted(set(run_counts) | set(wb_counts))
        table(
            ["code", "grp", "recorded run @327ee53", f"{s}_post.xlsx", "delta"],
            [[c, issue_group(c), run_counts.get(c, 0), wb_counts.get(c, 0),
              wb_counts.get(c, 0) - run_counts.get(c, 0)] for c in codes],
        )
        print(f"  totals: recorded run {sum(run_counts.values())}   "
              f"{s}_post.xlsx {sum(wb_counts.values())}")

        # the fields the completeness KPI is measured over
        by_id = {r.record_id: r for r in recs}
        rows_out = []
        for f in KPI_FIELDS:
            a = sum(1 for r in recs if not is_blank(getattr(r, f)))
            b = sum(1 for r in post.records if not is_blank(getattr(r, f)))
            if a or b:
                rows_out.append([sap_label(f), a, b, b - a])
        print()
        table([f"field", "filled, recorded run", f"filled, {s}_post.xlsx",
               "delta"], rows_out)


# ── main ─────────────────────────────────────────────────────────────────

SECTIONS = {
    "pairing": section_pairing,
    "vocabulary": section_vocabulary,
    "issues": section_issues,
    "divergence": section_divergence,
    "completeness": section_completeness,
    "improvement": section_improvement,
    "clustering": section_clustering,
    "cost": section_cost,
    "determinism": section_determinism,
    "recorded": section_recorded,
}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--section", choices=sorted(SECTIONS), action="append")
    args = ap.parse_args(argv)
    wanted = args.section or list(SECTIONS)

    resolved = pairs()
    audits: dict[str, dict[str, Audit]] = {}
    need_audits = {"pairing", "issues", "divergence", "completeness",
                   "improvement", "clustering", "recorded"}
    if need_audits & set(wanted):
        for s, (pre, post) in resolved.items():
            audits[s] = {"pre": Audit(pre), "post": Audit(post)}

    for name in SECTIONS:
        if name not in wanted:
            continue
        fn = SECTIONS[name]
        if name in ("vocabulary", "cost", "determinism"):
            fn()
        else:
            fn(audits)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
