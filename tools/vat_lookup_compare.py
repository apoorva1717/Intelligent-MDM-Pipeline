"""Compare the two missing-VAT lookup modes against known VAT IDs.

Usage::

    python tools/vat_lookup_compare.py tests/fixtures/vat/VAT_TEST_EXPORT_30.xlsx \\
        --rows 2-18 --exclude 9 --cache-dir <short path> [--out compare.xlsx] [--frozen]

    python tools/vat_lookup_compare.py <endpoint result>.xlsx --score-existing \\
        [--truth-column "VAT Actual"]

``--score-existing`` runs no lookup and no VIES check: it scores a workbook
the lookup endpoint already wrote (``VAT_Found_*``, plus ``VAT_Grounded_*``
when the LLM-only experiment ran next to the grounded lookup), against the
SAP VAT ID in ``--truth-column``.

Takes an export where some rows already have a VAT ID. For the chosen rows
(``--rows`` / ``--exclude`` as sheet row numbers; default: every row with a
value and Country/Region Key DE) it **blanks** ``VAT Registration No.``, keeps
the original as ground truth, and runs both modes on the blanked file:

* Mode A — ``grounded`` (``vat/lookup_grounded.py``);
* Mode W — ``llm_web`` (``vat/lookup_web.py``): the model searches the web,
  the code checks the page it cites; scored like Mode A (tiers included).
  Pass ``--modes grounded,llm_web``.
* Mode B — ``llm_only`` (``vat/lookup_llm_only.py``): **an experiment, not for
  production use; its answers are never written to SAP.**

Rows that were empty to begin with (a genuinely missing ID) are looked up too
and listed without a ground truth.

Output — a copy of the workbook (the Data sheet untouched) plus:

* ``Comparison`` — per row: ground truth, Mode A value / status / attribution
  verdict / candidates read / outcome, Mode B value / status / confidence /
  outcome; totals per mode, and Mode B's correctness by confidence;
* ``Cost`` — SerpAPI queries, pages fetched, LLM calls, VIES checks, cache
  write failures, per mode;
* ``Mode A rows`` / ``Mode B rows`` — the full ``VAT_Found_*`` output.

Outcomes, per value proposed:

* ``CORRECT`` — the same VAT ID as the ground truth (normalised: prefix +
  digits; spacing and case ignored);
* ``WRONG_VIES_VALID`` — wrong, passes the format check, and VIES says VALID:
  the worst case, a real number that belongs to someone else;
* ``WRONG_PLAUSIBLE`` — wrong, passes the format check, VIES not VALID;
* ``WRONG_FORMAT`` — wrong and fails the check digit / format;
* ``NULL`` — no value proposed.

Every wrong value that passes the format check gets a VIES answer: from the
mode itself when it asked, otherwise from an audit check made here — Mode A
does not ask VIES once a gate has failed, and without this the worst-case
count of the two modes would not be comparable.

A grounded value counts as a **proposal** only when every gate passed and its
attribution is ``ATTRIBUTION_MATCH`` or ``ATTRIBUTION_QUALIFIER`` (status
``FOUND_VERIFIED`` or ``FOUND_UNVERIFIED:VIES_*``). A value behind a failed
gate (``ATTRIBUTION_MISMATCH`` / ``_PARTIAL``, …), ``NOT_FOUND`` and
``LOOKUP_INCOMPLETE`` are "no proposal". Mode A's headline totals are
Proposed / Correct / Wrong / No proposal, with Precision = Correct / Proposed
and Recall = Correct / Customers with an SAP VAT ID; the old count — any value
returned, gate passed or not — is kept as a separate line.

When the grounded lookup recorded a tier (``VAT_Found_Tier`` /
``VAT_Grounded_Tier``), precision and recall are also given per tier
(PROPOSED, LIKELY, REVIEW) and cumulative (PROPOSED; PROPOSED + LIKELY). The
attribution-only proposal count above is kept next to it, to compare with
runs scored before tiers existed.

Totals count each Customer once (a customer split over several rows is one
lookup).
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import os
import re
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any, Callable

_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT))

logger = logging.getLogger("tools.vat_lookup_compare")

CORRECT = "CORRECT"
WRONG_VIES_VALID = "WRONG_VIES_VALID"
WRONG_PLAUSIBLE = "WRONG_PLAUSIBLE"
WRONG_FORMAT = "WRONG_FORMAT"
NULL = "NULL"
NO_TRUTH = "NO_TRUTH"
OUTCOMES = (CORRECT, WRONG_VIES_VALID, WRONG_PLAUSIBLE, WRONG_FORMAT, NULL)
CONFIDENCE_BUCKETS = ("high", "medium", "low", "null")
#: Modes whose results carry an attribution verdict and a tier.
TIERED_MODES = ("grounded", "llm_web")
#: Comparison-sheet column prefix per mode.
_PREFIX = {"grounded": "A", "llm_web": "W"}
#: Attribution verdicts under which a grounded value is a proposal.
PROPOSAL_ATTRIBUTIONS = ("ATTRIBUTION_MATCH", "ATTRIBUTION_QUALIFIER")


def parse_rows(spec: str | None) -> set[int]:
    """``"2-18,20"`` → {2, …, 18, 20}."""
    out: set[int] = set()
    for part in (spec or "").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            out.update(range(int(a), int(b) + 1))
        else:
            out.add(int(part))
    return out


def canon(value: str | None, country: str | None = "DE") -> str | None:
    """Comparison form of a VAT ID: normalised when it can be, else the
    compact upper-case characters."""
    from vat.normalize import assess_format

    if not value:
        return None
    fmt = assess_format(value, country)
    return fmt.normalized or re.sub(r"[^0-9A-Z]", "", str(value).upper()) or None


def _vies_from_status(status: str | None) -> str | None:
    """The VIES status a mode itself recorded, or None when it did not ask."""
    if not status:
        return None
    if status == "FOUND_VERIFIED":
        return "VALID"
    m = re.match(r"FOUND_UNVERIFIED:VIES_(\w+)$", status)
    if m:
        return m.group(1)
    m = re.match(r"LLM_GUESS:FORMAT_OK:(\w+)$", status)
    if m:
        return m.group(1)
    return None


def _first_flag(flags: str | None, prefix: str) -> str | None:
    for f in (flags or "").split("; "):
        if f.startswith(prefix):
            return f
    return None


def _count_flags(flags: str | None, prefix: str) -> int:
    return sum(1 for f in (flags or "").split("; ") if f.startswith(prefix))


def is_proposal(status: str | None, flags: str | None) -> bool:
    """A grounded value is a proposal only when every gate passed (status
    ``FOUND_VERIFIED`` or ``FOUND_UNVERIFIED:VIES_*``) and its attribution is
    MATCH or QUALIFIER. The winning result's own flags come before the
    ``CANDIDATE:`` entries, so the first ``ATTRIBUTION_*`` flag is its verdict."""
    if not status or not (status == "FOUND_VERIFIED" or status.startswith("FOUND_UNVERIFIED:VIES_")):
        return False
    return _first_flag(flags, "ATTRIBUTION_") in PROPOSAL_ATTRIBUTIONS


class _Audit:
    """The audit VIES checks made here, counted for the Cost sheet."""

    def __init__(self, client: Any) -> None:
        self.client = client
        self.checks = 0

    async def check(self, fmt: Any) -> Any:
        self.checks += 1
        return await self.client.check(fmt)


async def classify(value: str | None, truth: str | None, mode_vies: str | None,
                   audit: Any) -> tuple[str, str | None, bool | None]:
    """``(outcome, vies_status_used, passes_format)`` for one proposed value."""
    from vat.normalize import FORMAT_OK, assess_format

    if truth is None:  # a genuinely missing ID: nothing to be right or wrong about
        passes = assess_format(value, "DE").status == FORMAT_OK if value else None
        return NO_TRUTH, mode_vies, passes
    if not value:
        return NULL, None, None
    fmt = assess_format(value, "DE")
    passes = fmt.status == FORMAT_OK
    if canon(value) == truth:
        return CORRECT, mode_vies, passes
    if not passes:
        return WRONG_FORMAT, None, False
    vies = mode_vies
    if vies is None and audit is not None:
        vies = (await audit.check(fmt)).status
    return (WRONG_VIES_VALID if vies == "VALID" else WRONG_PLAUSIBLE), vies, True


async def compare(
    wb, settings: Any, *, rows: set[int] | None = None, exclude: set[int] | None = None,
    modes: tuple[str, ...] = ("grounded", "llm_only"),
    make_client: Callable[[], Any] | None = None,
    make_engine: Callable[[str, Any, Any], Any] | None = None,
) -> dict[str, Any]:
    """Run the comparison on *wb* in place (adds the result sheets). Returns
    the totals and costs. *make_client* / *make_engine* are injectable."""
    from vat.lookup_engines import make_lookup
    from vat.validate import resolve_columns, validate_rows
    from vat.vies_client import VIESClient
    from vat.xlsx import _data_sheet, _read
    from vat.lookup import LOOKUP_COLUMNS
    from vat.validate import VAT_COLUMNS

    make_client = make_client or (lambda: VIESClient(settings))
    make_engine = make_engine or make_lookup

    ws = _data_sheet(wb)
    _, headers, indexed = _read(ws, VAT_COLUMNS + LOOKUP_COLUMNS)
    cols = resolve_columns(headers)
    vat_col, country_col, customer_col = cols["vat"], cols["country"], cols["customer"]
    name_col = cols["name_1"]
    city_col = resolve_columns(headers, {"city": "city"})["city"]

    def cell(row, col):
        v = row.get(col) if col else None
        return str(v).strip() if v is not None and str(v).strip() else None

    # Ground truth: the chosen rows' values, then blank them.
    truth_by_row: dict[int, str] = {}
    blanked: list[dict[str, Any]] = []
    for sheet_row, row in indexed:
        value, country = cell(row, vat_col), cell(row, country_col)
        chosen = (sheet_row in rows) if rows else (value is not None and country == "DE")
        if exclude and sheet_row in exclude:
            chosen = False
        row = dict(row)
        if chosen and value:
            truth_by_row[sheet_row] = canon(value, country)
            row[vat_col] = None
        blanked.append(row)

    results: dict[str, list[dict[str, Any]]] = {}
    costs: dict[str, dict[str, Any]] = {}
    for mode in modes:
        client = make_client()
        engine = make_engine(mode, settings, client)
        started = time.perf_counter()
        try:
            out, summary = await validate_rows(blanked, client, lookup=engine)
        finally:
            await engine.aclose()
        results[mode] = out
        costs[mode] = {**summary["lookup"]["cost"],
                       "seconds": round(time.perf_counter() - started, 1)}

    audit = _Audit(make_client())
    table: list[dict[str, Any]] = []
    for (sheet_row, row), *mode_rows in zip(indexed, *(results[m] for m in modes)):
        if not any(r.get("VAT_Found_Status") for r in mode_rows):
            continue  # out of scope: not blank, or not DE
        truth = truth_by_row.get(sheet_row)
        entry: dict[str, Any] = {
            "sheet_row": sheet_row, "customer": cell(row, customer_col),
            "name": cell(row, name_col), "city": cell(row, city_col), "truth": truth,
        }
        for mode, r in zip(modes, mode_rows):
            entry[mode] = await _mode_entry(mode, r, truth, audit)
        table.append(entry)
    totals = _totals(table, modes)
    confidence = _by_confidence(table) if "llm_only" in modes else {}
    costs["audit"] = {"vies_checks": audit.checks,
                      "cache_write_failures": audit.client.write_failures.count}
    _write_sheets(wb, table, totals, confidence, costs, modes, results, indexed, cols)
    return {"rows": table, "totals": totals, "confidence": confidence, "costs": costs}


async def _mode_entry(mode: str, r: dict[str, Any], truth: str | None, audit: Any) -> dict[str, Any]:
    """One mode's result on one row, from its ``VAT_Found_*`` columns."""
    value, status, flags = r.get("VAT_Found_Value"), r.get("VAT_Found_Status"), r.get("VAT_Found_Flags")
    outcome, vies, passes = await classify(value, truth, _vies_from_status(status), audit)
    return {
        "value": value, "status": status, "outcome": outcome, "vies": vies,
        "passes_format": passes, "flags": flags,
        "attribution": _first_flag(flags, "ATTRIBUTION_"),
        "candidates": _count_flags(flags, "CANDIDATE:"),
        "confidence": r.get("VAT_Found_Evidence") if mode == "llm_only" else None,
        "source_url": r.get("VAT_Found_Source_URL"),
        "proposal": is_proposal(status, flags) if mode in TIERED_MODES else None,
        "tier": r.get("VAT_Found_Tier") if mode in TIERED_MODES else None,
    }


async def score_results(wb, *, truth_column: str = "VAT Actual") -> dict[str, Any]:
    """Score a workbook the lookup endpoint already wrote, in place (adds the
    result sheets). No lookup and no VIES check runs: a wrong value's VIES
    status is the one the mode recorded, if any.

    The grounded result is read from ``VAT_Grounded_*`` when that column set
    is filled (the LLM-only experiment ran next to it; ``VAT_Found_*`` is then
    Mode B), else from ``VAT_Found_*``."""
    from vat.lookup import LOOKUP_COLUMNS
    from vat.validate import resolve_columns
    from vat.xlsx import _data_sheet, _read

    ws = _data_sheet(wb)
    _, headers, indexed = _read(ws, [])
    if truth_column not in headers:
        raise ValueError(f"No column {truth_column!r} in sheet {ws.title!r}.")
    cols = resolve_columns(headers)
    city_col = resolve_columns(headers, {"city": "city"})["city"]

    def cell(row, col):
        v = row.get(col) if col else None
        return str(v).strip() if v is not None and str(v).strip() else None

    side_by_side = any(cell(row, "VAT_Grounded_Status") for _, row in indexed)
    # VAT_Found_* holds whichever mode VAT_Lookup_Mode names.
    labels = {cell(row, "VAT_Lookup_Mode") for _, row in indexed} - {None}
    found_mode = ("llm_web" if "LLM_WEB" in labels else "llm_only" if "LLM_ONLY" in labels
                  else "grounded")
    sources = ({"grounded": "VAT_Grounded_", found_mode: "VAT_Found_"} if side_by_side
               else {found_mode: "VAT_Found_"})
    modes = tuple(sources)
    results = {m: [{c: row.get(c.replace("VAT_Found_", prefix)) for c in LOOKUP_COLUMNS}
                   for _, row in indexed]
               for m, prefix in sources.items()}

    table: list[dict[str, Any]] = []
    for i, (sheet_row, row) in enumerate(indexed):
        if not any(results[m][i].get("VAT_Found_Status") for m in modes):
            continue  # out of scope
        raw_truth = cell(row, truth_column)
        entry: dict[str, Any] = {
            "sheet_row": sheet_row, "customer": cell(row, cols["customer"]),
            "name": cell(row, cols["name_1"]), "city": cell(row, city_col),
            "truth": canon(raw_truth, cell(row, cols["country"])) if raw_truth else None,
        }
        for m in modes:
            entry[m] = await _mode_entry(m, results[m][i], entry["truth"], None)
        table.append(entry)
    totals = _totals(table, modes)
    confidence = _by_confidence(table) if "llm_only" in modes else {}
    costs: dict[str, dict[str, Any]] = {"scored": {"rows": len(table), "truth_column": truth_column,
                                                   "lookups_run": 0, "vies_checks": 0}}
    _write_sheets(wb, table, totals, confidence, costs, modes, results, indexed, cols,
                  truth_note=f"Ground truth is the SAP VAT ID in {truth_column!r}; scored from an "
                             "existing result file — no lookup or VIES check was run.")
    return {"rows": table, "totals": totals, "confidence": confidence, "costs": costs}


def _per_customer(table: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[Any] = set()
    out = []
    for e in table:
        key = e["customer"] or ("row", e["sheet_row"])
        if key in seen:
            continue
        seen.add(key)
        out.append(e)
    return out


def _totals(table, modes) -> dict[str, dict[str, int]]:
    with_truth = [e for e in _per_customer(table) if e["truth"]]
    totals: dict[str, dict[str, int]] = {}
    for mode in modes:
        c = Counter(e[mode]["outcome"] for e in with_truth)
        t = {o: c.get(o, 0) for o in OUTCOMES}
        t["customers"] = len(with_truth)
        t["wrong_passes_format"] = t[WRONG_VIES_VALID] + t[WRONG_PLAUSIBLE]
        if mode in TIERED_MODES:
            proposed = [e for e in with_truth if e[mode]["proposal"]]
            t["proposed"] = len(proposed)
            t["proposed_correct"] = sum(e[mode]["outcome"] == CORRECT for e in proposed)
            t["proposed_wrong"] = t["proposed"] - t["proposed_correct"]
            t["no_proposal"] = t["customers"] - t["proposed"]
            t["precision"] = round(t["proposed_correct"] / t["proposed"], 3) if t["proposed"] else None
            t["recall"] = round(t["proposed_correct"] / t["customers"], 3) if t["customers"] else None
            t["any_value"] = sum(bool(e[mode]["value"]) for e in with_truth)
            t["tiers"] = _tier_metrics(with_truth, t["customers"], mode)
            t["verified"] = sum(e[mode]["status"] == "FOUND_VERIFIED" for e in with_truth)
            t["verified_but_wrong"] = sum(
                e[mode]["status"] == "FOUND_VERIFIED" and e[mode]["outcome"] != CORRECT
                for e in with_truth)
        totals[mode] = t
    return totals


TIER_ROWS = (("PROPOSED", ("PROPOSED",)), ("LIKELY", ("LIKELY",)), ("REVIEW", ("REVIEW",)),
             ("PROPOSED + LIKELY", ("PROPOSED", "LIKELY")))


def _tier_metrics(with_truth, customers: int, mode: str = "grounded") -> dict[str, dict[str, Any]] | None:
    """Per tier and cumulative: customers in it, correct, wrong, precision,
    recall (Correct / Customers with an SAP VAT ID). None when no row has a
    tier (a file written before tiers existed)."""
    if not any(e[mode]["tier"] for e in with_truth):
        return None
    out: dict[str, dict[str, Any]] = {}
    for label, tiers in TIER_ROWS:
        rows = [e for e in with_truth if e[mode]["tier"] in tiers]
        correct = sum(e[mode]["outcome"] == CORRECT for e in rows)
        out[label] = {"n": len(rows), "correct": correct, "wrong": len(rows) - correct,
                      "precision": round(correct / len(rows), 3) if rows else None,
                      "recall": round(correct / customers, 3) if customers else None}
    return out


def _by_confidence(table) -> dict[str, dict[str, int]]:
    out: dict[str, dict[str, int]] = {b: Counter() for b in CONFIDENCE_BUCKETS}
    for e in _per_customer(table):
        if not e["truth"]:
            continue
        b = e["llm_only"]["confidence"]
        b = b if b in ("high", "medium", "low") else "null"
        out[b]["customers"] += 1
        out[b][e["llm_only"]["outcome"]] += 1
    return {b: {"customers": c.get("customers", 0), **{o: c.get(o, 0) for o in OUTCOMES}}
            for b, c in out.items()}


_IMPRESSUM_TRUTH_NOTE = ("Ground truth came from the companies' Impressum pages, so Mode A recall is not an "
                         "independent measure.")


def _write_sheets(wb, table, totals, confidence, costs, modes, results, indexed, cols,
                  truth_note: str = _IMPRESSUM_TRUTH_NOTE) -> None:
    from vat.lookup import LOOKUP_COLUMNS

    for title in ("Comparison", "Cost", "Mode A rows", "Mode W rows", "Mode B rows"):
        if title in wb.sheetnames:
            del wb[title]

    ws = wb.create_sheet("Comparison")
    ws.append(["Missing-VAT lookup comparison. Proposals only — nothing here is written to SAP."])
    if "llm_only" in modes:
        ws.append(["Mode B (llm_only) is an EXPERIMENT, not for production use."])
    ws.append([truth_note])
    ws.append([])
    header = ["Sheet row", "Customer", "Name", "City", "Ground truth"]
    for m in (m for m in TIERED_MODES if m in modes):
        x = _PREFIX[m]
        header += [f"{x} value", f"{x} status", f"{x} attribution", f"{x} candidates read", f"{x} VIES",
                   f"{x} outcome", f"{x} correct?", f"{x} proposal?", f"{x} tier"]
    if "llm_only" in modes:
        header += ["B value", "B status", "B confidence", "B check digit ok", "B VIES", "B outcome",
                   "B correct?"]
    ws.append(header)
    for e in table:
        line = [e["sheet_row"], e["customer"], e["name"], e["city"], e["truth"] or "(none — genuinely missing)"]
        for m in (m for m in TIERED_MODES if m in modes):
            a = e[m]
            line += [a["value"], a["status"], a["attribution"], a["candidates"], a["vies"],
                     a["outcome"], _yes_no(a["outcome"]), _tri(a["proposal"]), a["tier"]]
        if "llm_only" in modes:
            b = e["llm_only"]
            line += [b["value"], b["status"], b["confidence"], _tri(b["passes_format"]), b["vies"],
                     b["outcome"], _yes_no(b["outcome"])]
        ws.append(line)

    ws.append([])
    ws.append(["Totals (each Customer counted once; rows without a ground truth excluded)"])
    ws.append(["Outcome", *[_mode_label(m) for m in modes]])
    labels = [
        ("customers", "Customers with a ground truth"),
        (CORRECT, "Correct"),
        ("wrong_passes_format", "Wrong but plausible (passes format)"),
        (WRONG_VIES_VALID, "  of which VIES VALID (worst case: someone else's real number)"),
        (WRONG_PLAUSIBLE, "  of which VIES not VALID"),
        (WRONG_FORMAT, "Wrong, fails check digit / format"),
        (NULL, "Null (no value proposed)"),
    ]
    for key, label in labels:
        ws.append([label, *[totals[m].get(key, 0) for m in modes]])
    for m in (m for m in TIERED_MODES if m in modes):
        x, a = f"Mode {_PREFIX[m]}", totals[m]
        ws.append([f"{x} FOUND_VERIFIED", a["verified"]])
        ws.append([f"{x} FOUND_VERIFIED but wrong", a["verified_but_wrong"]])
        ws.append([])
        ws.append([f"{x} proposals (gates passed, attribution MATCH or QUALIFIER; per Customer)"])
        for label, key in (("Customers with an SAP VAT ID", "customers"), ("Proposed", "proposed"),
                           ("Correct", "proposed_correct"), ("Wrong", "proposed_wrong"),
                           ("No proposal", "no_proposal"),
                           ("Precision (Correct / Proposed)", "precision"),
                           ("Recall (Correct / Customers with an SAP VAT ID)", "recall")):
            ws.append([label, a[key]])
        ws.append(["Old metric: any value returned (gate passed or not)", a["any_value"],
                   "of which correct", a[CORRECT]])
        if a.get("tiers"):
            ws.append([])
            ws.append([f"{x} by tier (per Customer; recall = Correct / Customers with an SAP VAT ID)"])
            ws.append(["Tier", "Customers", "Correct", "Wrong", "Precision", "Recall"])
            for label, m in a["tiers"].items():
                ws.append([label, m["n"], m["correct"], m["wrong"], m["precision"], m["recall"]])

    if confidence:
        ws.append([])
        ws.append(["Mode B correctness by self-reported confidence (per Customer)"])
        ws.append(["Confidence", "Customers", "Correct", "Wrong, VIES VALID", "Wrong, plausible",
                   "Wrong, bad format", "Null"])
        for b in CONFIDENCE_BUCKETS:
            c = confidence[b]
            ws.append([b, c["customers"], c[CORRECT], c[WRONG_VIES_VALID], c[WRONG_PLAUSIBLE],
                       c[WRONG_FORMAT], c[NULL]])

    cost = wb.create_sheet("Cost")
    keys = sorted({k for m in costs for k in costs[m]})
    cost.append(["Metric", *[_mode_label(m) for m in costs]])
    for k in keys:
        cost.append([k, *[costs[m].get(k, "") for m in costs]])

    for mode, title in (("grounded", "Mode A rows"), ("llm_web", "Mode W rows"),
                        ("llm_only", "Mode B rows")):
        if mode not in results:
            continue
        sheet = wb.create_sheet(title)
        sheet.append(["Sheet row", "Customer", "Name", *LOOKUP_COLUMNS])
        for (sheet_row, row), r in zip(indexed, results[mode]):
            if r.get("VAT_Found_Status"):
                sheet.append([sheet_row, row.get(cols["customer"]), row.get(cols["name_1"]),
                              *[r.get(c) for c in LOOKUP_COLUMNS]])


def _mode_label(mode: str) -> str:
    return {"grounded": "Mode A (grounded)", "llm_web": "Mode W (LLM web search)",
            "llm_only": "Mode B (LLM only, experiment)",
            "audit": "Audit VIES checks"}.get(mode, mode)


def _yes_no(outcome: str) -> str:
    return {CORRECT: "yes", NO_TRUTH: "n/a", NULL: "null"}.get(outcome, "no")


def _tri(value: bool | None) -> str | None:
    return None if value is None else ("yes" if value else "no")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="Export XLSX with known VAT IDs (sheet 'Data').")
    ap.add_argument("--out", help="Output XLSX (default: <input>_lookup_compare.xlsx).")
    ap.add_argument("--rows", help="Sheet rows whose VAT ID is the ground truth, e.g. 2-18.")
    ap.add_argument("--exclude", help="Sheet rows to leave out of --rows, e.g. 9.")
    ap.add_argument("--modes", default="grounded,llm_only")
    ap.add_argument("--frozen", action="store_true", help="CACHE_FROZEN: replay recordings only.")
    ap.add_argument("--cache-dir", default=None, help="Override EVIDENCE_CACHE_DIR.")
    ap.add_argument("--score-existing", action="store_true",
                    help="Score an existing endpoint result file; no lookup or VIES check runs.")
    ap.add_argument("--truth-column", default="VAT Actual",
                    help="With --score-existing: the column holding the SAP VAT ID.")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    if args.frozen:
        os.environ["CACHE_FROZEN"] = "true"
    if args.cache_dir is not None:
        os.environ["EVIDENCE_CACHE_DIR"] = args.cache_dir

    from openpyxl import load_workbook

    from config import get_settings

    input_path = Path(args.input).resolve()
    out_path = Path(args.out).resolve() if args.out else input_path.with_name(
        f"{input_path.stem}_lookup_compare{input_path.suffix}")
    if out_path == input_path:
        raise SystemExit("--out must differ from the input file.")
    wb = load_workbook(input_path)
    if args.score_existing:
        report = asyncio.run(score_results(wb, truth_column=args.truth_column))
        wb.save(out_path)
        print(f"Wrote {out_path}")
        _print_proposals(report["totals"])
        return 0
    report = asyncio.run(compare(
        wb, get_settings(), rows=parse_rows(args.rows) or None,
        exclude=parse_rows(args.exclude) or None,
        modes=tuple(m.strip() for m in args.modes.split(",") if m.strip()),
    ))
    wb.save(out_path)
    print(f"Wrote {out_path}")
    for mode, t in report["totals"].items():
        print(f"  {_mode_label(mode)}: {t}")
    for mode, c in report["costs"].items():
        print(f"  cost {_mode_label(mode)}: {c}")
    _print_proposals(report["totals"])
    return 0


def _print_proposals(totals: dict[str, dict[str, Any]]) -> None:
    for mode in TIERED_MODES:
        if mode in totals:
            _print_mode_proposals(f"Mode {_PREFIX[mode]}", totals[mode])


def _print_mode_proposals(x: str, a: dict[str, Any]) -> None:
    print(f"  {x}: customers with SAP VAT {a['customers']} | proposed {a['proposed']} | "
          f"correct {a['proposed_correct']} | wrong {a['proposed_wrong']} | "
          f"no proposal {a['no_proposal']} | precision {a['precision']} | recall {a['recall']}")
    print(f"  {x} old metric (any value): {a['any_value']} returned, {a[CORRECT]} correct, "
          f"{a['any_value'] - a[CORRECT]} wrong, {a[NULL]} null")
    for label, m in (a.get("tiers") or {}).items():
        print(f"  {x} tier {label}: {m['n']} | correct {m['correct']} | wrong {m['wrong']} | "
              f"precision {m['precision']} | recall {m['recall']}")


if __name__ == "__main__":
    sys.exit(main())
